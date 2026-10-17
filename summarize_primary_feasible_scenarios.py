#!/usr/bin/env python3
"""Describe the reference primary feasible set and reproduce Supplementary Table S4.

This is descriptive post-processing, not a new simulation or optimization.
The input is the original inspiration-only, volume-reset primary CSV (324 rows),
NOT outputs of the separate passive-expiration sensitivity analysis.

Run beside the original CSV:
    python summarize_primary_feasible_scenarios.py --output-dir .
Use a different input location:
    python summarize_primary_feasible_scenarios.py --input-csv path/to/MVtarget_primary_feasible_simulation_results.csv --output-dir results
Omit Word output:
    python summarize_primary_feasible_scenarios.py --output-dir . --no-docx

Dependencies: NumPy; python-docx only for the optional Word table.
All input files remain unchanged. Counts and percentiles describe the prespecified
simulation grid; no inferential tests, clinical prevalence estimates, or confidence
intervals are computed. No interpolation of model responses or settings is done.
"""
from __future__ import annotations

import argparse
from collections import Counter
import csv
import hashlib
import itertools
import json
from pathlib import Path
import sys
from typing import Any

import numpy as np

ORDER = ("compliance_dominant", "resistance_dominant", "mixed", "severe_mixed")
LABELS = {
    "compliance_dominant": "Compliance-dominant",
    "resistance_dominant": "Resistance-dominant",
    "mixed": "Mixed",
    "severe_mixed": "Severe mixed",
}
KEYS = ("phenotype", "VT_ml_per_kg", "RR_min", "Ti_s", "waveform", "pause_fraction", "PEEP_cmH2O")
NUMERIC = (
    "PBW_kg", "VT_ml_per_kg", "VT_ml", "RR_min", "Ti_s", "pause_fraction",
    "PEEP_cmH2O", "minute_ventilation_L_min", "MP_J_min", "EII", "DCES",
    "Ppeak_cmH2O", "end_inspiratory_volume_L",
)
METRICS = (("MP_J_min", "MP", "J/min", 2), ("EII", "EII", "dimensionless", 3),
           ("DCES", "DCES", "dimensionless", 3))
PAIRS = ((4.0, 25.0), (5.0, 20.0), (7.0, 15.0))
FLOW_ORDER = ("square", "decelerating", "sinusoidal")
FLOAT_ATOL = 1e-10


def read_csv(path: Path) -> list[dict[str, Any]]:
    """Read numerical CSV values without modifying the source."""
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        required = set(NUMERIC) | {"phenotype", "waveform"}
        missing = required - set(reader.fieldnames or ())
        if missing:
            raise ValueError(f"Missing columns in {path.name}: {sorted(missing)}")
        rows = list(reader)
    if not rows:
        raise ValueError(f"Empty dataset: {path}")
    for index, row in enumerate(rows, start=2):
        for column in NUMERIC:
            try:
                value = float(row[column])
            except (TypeError, ValueError) as exc:
                raise ValueError(f"Invalid {column} in CSV row {index}") from exc
            if not np.isfinite(value):
                raise ValueError(f"Non-finite {column} in CSV row {index}")
            row[column] = value
    return rows


def scenario_key(row: dict[str, Any]) -> tuple[Any, ...]:
    return tuple(row[key] for key in KEYS)


def validate_primary(rows: list[dict[str, Any]]) -> None:
    """Check the reference design and distinguish it from periodic-state output."""
    if len(rows) != 324 or Counter(r["phenotype"] for r in rows) != Counter(dict.fromkeys(ORDER, 81)):
        raise ValueError("Expected the reference primary set: 324 scenarios, 81 per configuration.")
    expected = {
        (p, vt, rr, ti, wave, pause, 5.0)
        for p in ORDER for vt, rr in PAIRS
        for ti, wave, pause in itertools.product((0.6, 1.0, 1.5), FLOW_ORDER, (0.0, 0.1, 0.2))
    }
    observed = [scenario_key(r) for r in rows]
    if len(set(observed)) != len(rows) or set(observed) != expected:
        raise ValueError("The retained setting combinations do not match the reference grid.")
    for r in rows:
        if not np.isclose(r["PBW_kg"], 70.0, rtol=0, atol=FLOAT_ATOL):
            raise ValueError("Expected predicted body weight 70 kg.")
        if not np.isclose(r["VT_ml"], r["VT_ml_per_kg"] * r["PBW_kg"], rtol=0, atol=FLOAT_ATOL):
            raise ValueError("Inconsistent tidal volume columns.")
        mv = r["VT_ml"] * r["RR_min"] / 1000.0
        if not np.isclose(r["minute_ventilation_L_min"], mv, rtol=0, atol=FLOAT_ATOL):
            raise ValueError("Inconsistent minute ventilation.")
        if not (6.5 - FLOAT_ATOL <= mv <= 7.5 + FLOAT_ATOL and r["Ppeak_cmH2O"] <= 35.0 + FLOAT_ATOL):
            raise ValueError("Scenario fails the primary feasibility criteria.")
        if not np.isclose(r["end_inspiratory_volume_L"], r["VT_ml"] / 1000.0, rtol=0, atol=FLOAT_ATOL):
            raise ValueError("Expected reference volume-reset data, not periodic-expiration output.")
        if r["MP_J_min"] <= 0 or r["EII"] < 0 or not (1.0 / 6.0 - FLOAT_ATOL <= r["DCES"] <= 1):
            raise ValueError("Invalid energy metric value.")
        if "analysis_name" in r and r["analysis_name"] != "primary_MV_target_7.0_L_min":
            raise ValueError("Unexpected analysis label; use the reference primary CSV.")


def verify_full_dataset(primary: list[dict[str, Any]], path: Path) -> dict[str, Any]:
    """Independently reproduce feasibility membership when the full CSV is supplied."""
    full = read_csv(path)
    if len(full) != 2700 or Counter(r["phenotype"] for r in full) != Counter(dict.fromkeys(ORDER, 675)):
        raise ValueError("Expected 2700 full-grid scenarios, 675 per configuration.")
    selected = [r for r in full if r["Ppeak_cmH2O"] <= 35.0
                and 6.5 <= r["minute_ventilation_L_min"] <= 7.5]
    lookup = {scenario_key(r): r for r in selected}
    if len(lookup) != 324 or set(lookup) != {scenario_key(r) for r in primary}:
        raise ValueError("Primary membership does not match filtering of the full-grid CSV.")
    largest = 0.0
    for r in primary:
        for metric in ("MP_J_min", "EII", "DCES", "Ppeak_cmH2O"):
            difference = abs(r[metric] - lookup[scenario_key(r)][metric])
            largest = max(largest, difference)
            if not np.isclose(r[metric], lookup[scenario_key(r)][metric], rtol=1e-12, atol=1e-12):
                raise ValueError(f"Reference output mismatch for {metric}.")
    return {"full_rows": len(full), "matched_primary_rows": len(lookup),
            "peak_pressure_exclusions_in_full_grid": sum(r["Ppeak_cmH2O"] > 35.0 for r in full),
            "maximum_absolute_metric_difference": largest,
            "full_csv_sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def summaries(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    metric_rows: list[dict[str, Any]] = []
    count_rows: list[dict[str, Any]] = []
    for p in ORDER:
        selected = [r for r in rows if r["phenotype"] == p]
        count_rows.append({"configuration": "all", "variable": "configuration",
                           "level": p, "n": len(selected), "denominator": len(rows),
                           "percent": 100.0 * len(selected) / len(rows)})
        for key, label, unit, decimals in METRICS:
            values = np.array([r[key] for r in selected], dtype=np.float64)
            # NumPy linear percentiles; with n=81, Q1, median and Q3 all fall
            # exactly on observed order statistics, not between scenarios.
            low, q1, median, q3, high = np.quantile(values, [0, .25, .5, .75, 1], method="linear")
            metric_rows.append({"configuration": p, "metric": label, "unit": unit,
                                "n": len(selected), "minimum": float(low),
                                "p25": float(q1), "median": float(median),
                                "p75": float(q3), "maximum": float(high)})
    for p in (*ORDER, "all"):
        selected = rows if p == "all" else [r for r in rows if r["phenotype"] == p]
        for variable in ("VT_ml_per_kg", "RR_min", "VT_RR_pair", "Ti_s", "waveform",
                         "pause_fraction", "minute_ventilation_L_min", "PEEP_cmH2O"):
            labels = []
            for r in selected:
                if variable == "VT_RR_pair":
                    label = f"{r['VT_ml_per_kg']:g} mL/kg; {r['RR_min']:g} breaths/min"
                elif variable == "waveform":
                    label = str(r[variable])
                elif variable == "minute_ventilation_L_min":
                    # Merge only floating-point representations of identical MV
                    # (e.g. 7 and 7.000000000000001) for categorical counting.
                    label = f"{r['VT_ml'] * r['RR_min'] / 1000.0:.2f}"
                else:
                    label = f"{r[variable]:g}"
                labels.append(label)
            for level, n in sorted(Counter(labels).items()):
                count_rows.append({"configuration": p, "variable": variable, "level": level,
                                   "n": n, "denominator": len(selected),
                                   "percent": 100.0 * n / len(selected)})
    return metric_rows, count_rows


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def export_docx(rows: list[dict[str, Any]], metrics: list[dict[str, Any]], output: Path) -> Path:
    """Build the two-page, monochrome, editable supplementary table."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.45)
    sec.top_margin = sec.bottom_margin = Cm(1.25)
    sec.header_distance = sec.footer_distance = Cm(.6)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.04
    for name in ("Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(5)

    def para(text: str, *, bold: bool = False, size: float = 10, after: float = 6) -> Any:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.keep_with_next = bold
        run = p.add_run(text)
        run.bold, run.font.size = bold, Pt(size)
        return p

    def table(headers: list[str], cells: list[list[str]], widths: list[float], *, size: float = 9.5) -> Any:
        t = doc.add_table(rows=1, cols=len(headers))
        t.autofit = False
        for c, width in zip(t.columns, widths):
            c.width = Cm(width)
        for i, text in enumerate(headers):
            t.rows[0].cells[i].text = text
        repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for values in cells:
            for cell, text in zip(t.add_row().cells, values):
                cell.text = text
        for ridx, row in enumerate(t.rows):
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
            for cidx, (cell, width) in enumerate(zip(row.cells, widths)):
                cell.width = Cm(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                margins = OxmlElement("w:tcMar")
                for side, value in (("top", "90"), ("bottom", "90"), ("left", "100"), ("right", "100")):
                    element = OxmlElement("w:" + side)
                    element.set(qn("w:w"), value)
                    element.set(qn("w:type"), "dxa")
                    margins.append(element)
                tcpr.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.0
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if cidx == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.name, run.font.size = "Arial", Pt(size)
                        run.font.color.rgb = RGBColor(0, 0, 0)
                        run.bold = ridx == 0
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            element = OxmlElement("w:" + side)
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), "4")
            element.set(qn("w:color"), "000000")
            borders.append(element)
        t._tbl.tblPr.append(borders)
        return t

    para("Supplementary Table S4. Characterization of the reference primary feasible scenarios.", bold=True, size=12)
    para("All summaries below use the original inspiration-only, volume-reset dataset (324 scenarios). "
         "Outputs of the separate passive-expiration sensitivity analysis are not included.", size=9.5)
    doc.add_paragraph("A. Scenario distribution across mechanical configurations", style="Heading 2")
    table(["Configuration", "Feasible scenarios (n)", "Share of the primary set (%)"],
          [[LABELS[p], str(sum(r["phenotype"] == p for r in rows)), "25.0"] for p in ORDER] +
          [["Total", str(len(rows)), "100.0"]], [10.6, 8.1, 8.1])
    doc.add_paragraph("B. Retained ventilator settings and their frequencies", style="Heading 2")
    para("The same setting combinations were retained in all four configurations. Tidal volume and "
         "respiratory rate are shown jointly because their retained values are not freely interchangeable.", size=9.5)
    pair_cells = []
    for vt, rr in PAIRS:
        n = sum(r["VT_ml_per_kg"] == vt and r["RR_min"] == rr for r in rows)
        pair_cells.append([f"{vt:g}", f"{vt * 70:g}", f"{rr:g}", f"{vt * 70 * rr / 1000:.2f}",
                           f"{n // 4} (33.3%)", f"{n} (33.3%)"])
    table(["VT (mL/kg PBW)", "VT (mL)", "RR (breaths/min)", "MV (L/min)",
           "n (%) per configuration", "n (%) across all configurations"], pair_cells,
          [4.15, 3.0, 4.05, 3.2, 5.95, 6.45])
    para("Each retained VT–RR pair covered the complete 3 × 3 × 3 combination of inspiratory time, "
         "flow waveform, and pause fraction (27 scenarios per configuration).", size=9.5, after=6)
    table(["Parameter", "Retained levels", "n per level, per configuration", "n per level, entire primary set"],
          [["Inspiratory time (s)", "0.6; 1.0; 1.5", "27", "108"],
           ["Flow waveform", "Square; linearly decelerating; sinusoidal", "27", "108"],
           ["End-inspiratory pause fraction", "0.0; 0.1; 0.2", "27", "108"]],
          [6.0, 9.3, 5.75, 5.75])
    para("Counts in the last table are marginal counts for each individual level, not additional scenarios. "
         "MV was 7.00 L/min in 54/81 scenarios (66.7%) per configuration and 7.35 L/min in 27/81 (33.3%); "
         "the corresponding totals were 216 and 108 scenarios. PBW was 70 kg and PEEP was 5 cmH₂O throughout.", size=9.2)
    para("All retained scenarios satisfied Ppeak ≤35 cmH₂O. Frequencies reflect the predefined grid "
         "and feasibility filters, not the prevalence of patient phenotypes or clinical ventilator use.", size=9.2)

    doc.add_page_break()
    para("Supplementary Table S4 (continued).", bold=True, size=12)
    doc.add_paragraph("C. Distributions of global and regional energy metrics", style="Heading 2")
    para("Each distribution includes all 81 reference primary feasible scenarios in the stated configuration, "
         "rather than only the minimum-MP, balanced, or Pareto-selected solutions.", size=9.5)
    matrix = []
    for item in metrics:
        d = 2 if item["metric"] == "MP" else 3
        matrix.append([LABELS[item["configuration"]], "MP (J/min)" if d == 2 else item["metric"],
                       *[f"{item[k]:.{d}f}" for k in ("minimum", "p25", "median", "p75", "maximum")]])
    t = table(["Configuration (n = 81 each)", "Metric", "Minimum", "25th percentile", "Median", "75th percentile", "Maximum"],
              matrix, [6.2, 3.2, 3.35, 3.6, 3.35, 3.75, 3.35], size=10)
    # Merge configuration labels vertically within each three-metric block.
    for start in (1, 4, 7, 10):
        merged = t.cell(start, 0).merge(t.cell(start + 2, 0))
        merged.text = LABELS[metrics[start - 1]["configuration"]]
        merged.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = merged.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            run.font.name, run.font.size = "Arial", Pt(10)
    para("Percentiles were calculated from unrounded numerical outputs using the linear quantile convention. "
         "For each 81-scenario distribution, the 25th percentile, median, and 75th percentile coincide with "
         "the 21st, 41st, and 61st ordered observations, respectively. Values are rounded only for presentation.", size=9.5, after=8)
    para("The 25th–75th percentile interval describes variation across the evaluated scenarios; it is not "
         "a confidence interval. No inferential statistical comparisons were performed. The summarized "
         "regional metrics retain the original positive-inflow energy definition and are not validated "
         "measures of regional injury risk.", size=9.5, after=8)
    para("Abbreviations: VT, tidal volume; PBW, predicted body weight; RR, respiratory rate; MV, minute ventilation; "
         "PEEP, positive end-expiratory pressure; Ppeak, peak airway-opening pressure; MP, global inspiratory "
         "mechanical power above prescribed PEEP; EII, Energy Inequality Index; DCES, Dominant Compartment "
         "Energy Share. EII and DCES are dimensionless. Pause fraction denotes the fraction of total "
         "inspiratory time occupied by the end-inspiratory pause.", size=9.5)
    para("Numerical source: MVtarget_primary_feasible_simulation_results.csv. Descriptive summaries were "
         "generated without changing the model equations, scenario inclusion, optimization, or existing results.", size=9.2)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Supplementary Table S4 | ")
    r.font.name, r.font.size = "Arial", Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "Supplementary Table S4: Reference primary feasible scenarios"
    doc.core_properties.author = "Tomasz Urbankowski"
    path = output / "Supplementary_Table_S4_primary_feasible_scenarios.docx"
    doc.save(path)
    return path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input-csv", type=Path, default=Path("MVtarget_primary_feasible_simulation_results.csv"))
    parser.add_argument("--full-csv", type=Path, help="Optional original full-grid CSV for independent membership verification.")
    parser.add_argument("--output-dir", type=Path, default=Path("."))
    parser.add_argument("--no-docx", action="store_true", help="Write numerical summaries only.")
    args = parser.parse_args()
    try:
        source = args.input_csv.resolve()
        if not source.is_file():
            raise FileNotFoundError(f"Input CSV not found: {source}")
        out = args.output_dir.resolve()
        out.mkdir(parents=True, exist_ok=True)
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        rows = read_csv(source)
        validate_primary(rows)
        full_check = verify_full_dataset(rows, args.full_csv.resolve()) if args.full_csv else None
        metrics, counts = summaries(rows)
        paths = [out / "primary_feasible_metric_summary.csv", out / "primary_feasible_setting_counts.csv"]
        if source in paths:
            raise ValueError("Input and output paths must differ.")
        write_csv(paths[0], metrics)
        write_csv(paths[1], counts)
        if not args.no_docx:
            try:
                paths.append(export_docx(rows, metrics, out))
            except ImportError as exc:
                raise RuntimeError("Install python-docx or use --no-docx; numerical CSV summaries have been saved.") from exc
        if hashlib.sha256(source.read_bytes()).hexdigest() != digest:
            raise RuntimeError("Input file unexpectedly changed.")
        verification = {
            "analysis": "reference_primary_descriptive_characterization", "input_csv": source.name,
            "input_sha256": digest, "n": len(rows),
            "n_per_configuration": {p: sum(r["phenotype"] == p for r in rows) for p in ORDER},
            "quantile_method": "linear", "raw_values_used_for_percentiles": True,
            "original_input_unchanged": True, "simulations_rerun": False,
            "periodic_expiration_outputs_included": False, "inferential_tests_performed": False,
            "full_grid_cross_check": full_check, "postprocessing_python": sys.version.split()[0],
            "postprocessing_numpy": np.__version__,
        }
        check_path = out / "primary_feasible_summary_verification.json"
        check_path.write_text(json.dumps(verification, indent=2), encoding="utf-8")
        paths.append(check_path)
        print("Reference primary set verified: 324 scenarios, 81 per configuration.")
        for m in metrics:
            print(m["configuration"], m["metric"], m["median"], m["p25"], m["p75"])
        for path in paths:
            print(path)
    except (OSError, ValueError, RuntimeError) as exc:
        parser.exit(1, f"Error: {exc}\n")


if __name__ == "__main__":
    main()
