#!/usr/bin/env python3
"""Check prescribed-PEEP offsets in the reference linear, volume-reset model.

This is a structural pressure-offset/feasibility check, NOT a physiological PEEP
optimization. R_i, C_i, waveforms, relative initial volumes and energy definitions
are unchanged. PEEP-dependent recruitment, overdistension and absolute lung-volume
changes are not simulated. Neither the periodic-expiration model nor refined
parameter grids are used here.

Run beside the main model and reference CSV:
    python assess_peep_offset_sensitivity.py --output-dir . --no-docx

Default test values: prescribed PEEP 5, 10, 15, 20 cmH2O. These are illustrative
pressure offsets, not proposed clinical settings. Each of the original 324
primary setting combinations is directly re-evaluated at each PEEP. The script
reapplies Ppeak <= 35 cmH2O and repeats minimum-MP and balanced selection. It does
not recompute Pareto fronts or overwrite the original CSV/model.

Dependencies: numpy, pandas; python-docx only for optional Table S6 export.
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import itertools
import json
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

ORDER = ('compliance_dominant', 'resistance_dominant', 'mixed', 'severe_mixed')
LABELS = dict(zip(ORDER, ('Compliance-dominant', 'Resistance-dominant', 'Mixed', 'Severe mixed')))
PEEPS = (5.0, 10.0, 15.0, 20.0)
DT = 0.002
RTOL, ATOL = 1e-10, 1e-12
SETTINGS = ('VT_ml_per_kg', 'RR_min', 'Ti_s', 'waveform', 'pause_fraction')
ENERGETICS = ('global_energy_J', 'MP_J_min', 'EII', 'DCES')
REGIONAL = tuple(itertools.chain.from_iterable((f'REF_{i}', f'comp_energy_J_{i}') for i in range(1, 7)))


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_model(path: Path) -> Any:
    spec = importlib.util.spec_from_file_location('reference_ventilation_model', path)
    if spec is None or spec.loader is None:
        raise ValueError(f'Cannot import model: {path}')
    model = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(model)
    for name in ('get_phenotypes', 'simulate_breath'):
        if not callable(getattr(model, name, None)):
            raise ValueError(f'Model must provide {name}().')
    return model


def validate_input(x: pd.DataFrame) -> None:
    required = {'phenotype', 'PBW_kg', 'VT_ml', 'PEEP_cmH2O', 'minute_ventilation_L_min',
                'Ppeak_cmH2O', 'balanced_score', *SETTINGS, *ENERGETICS, *REGIONAL}
    if required - set(x.columns):
        raise ValueError(f'Missing columns: {sorted(required - set(x.columns))}')
    if x.shape[0] != 324 or x.groupby('phenotype').size().to_dict() != dict.fromkeys(ORDER, 81):
        raise ValueError('Expected the reference primary CSV: 324 scenarios, 81 per configuration.')
    numbers = list(required - {'phenotype', 'waveform'})
    if not np.isfinite(x[numbers].to_numpy(dtype=float)).all():
        raise ValueError('Input contains missing/non-finite numerical values.')
    if not (np.isclose(x.PEEP_cmH2O, 5.0, rtol=0, atol=ATOL).all()
            and np.isclose(x.PBW_kg, 70.0, rtol=0, atol=ATOL).all()):
        raise ValueError('Expected reference PEEP 5 cmH2O and PBW 70 kg.')
    if not np.allclose(x.VT_ml / 1000 * x.RR_min, x.minute_ventilation_L_min, rtol=RTOL, atol=ATOL):
        raise ValueError('Minute ventilation does not agree with VT and RR.')
    if not ((x.Ppeak_cmH2O <= 35) & x.minute_ventilation_L_min.between(6.5, 7.5)).all():
        raise ValueError('Input does not satisfy reference primary feasibility conditions.')
    expected = {(p, vt, rr, ti, wave, pause) for p in ORDER
                for vt, rr in ((4, 25), (5, 20), (7, 15))
                for ti, wave, pause in itertools.product((.6, 1., 1.5),
                                                       ('square', 'decelerating', 'sinusoidal'), (0., .1, .2))}
    observed = list(x[['phenotype', *SETTINGS]].itertuples(index=False, name=None))
    if len(set(observed)) != len(x) or set(observed) != expected:
        raise ValueError('Input does not match the original primary setting grid.')


def directly_replay(x: pd.DataFrame, model: Any) -> tuple[pd.DataFrame, dict[str, Any]]:
    mechanics = model.get_phenotypes()
    records: list[dict[str, Any]] = []
    metrics = [*ENERGETICS, *REGIONAL, 'end_inspiratory_volume_L']
    max_reference_difference = 0.0
    max_invariance_difference = 0.0
    max_pressure_offset_error = 0.0
    for sid, row in x.iterrows():
        reference_result = None
        for peep in PEEPS:
            result = model.simulate_breath(
                mechanics[row.phenotype]['C'], mechanics[row.phenotype]['R'],
                vt_l=float(row.VT_ml) / 1000., rr=float(row.RR_min), ti_s=float(row.Ti_s),
                waveform=str(row.waveform), pause_fraction=float(row.pause_fraction),
                peep_cmH2O=peep, dt=DT)
            values = np.asarray([result[c] for c in metrics], dtype=float)
            if not np.isfinite(values).all():
                raise ValueError(f'Non-finite model result at scenario {sid + 1}, PEEP {peep}.')
            if peep == 5.0:
                reference_result = result
                stored = np.asarray([row[c] for c in [*ENERGETICS, *REGIONAL, 'Ppeak_cmH2O']], dtype=float)
                replay = np.asarray([result[c] for c in [*ENERGETICS, *REGIONAL, 'Ppeak_cmH2O']], dtype=float)
                if not np.allclose(stored, replay, rtol=RTOL, atol=ATOL):
                    raise ValueError(f'Reference replay mismatch for scenario {sid + 1}.')
                max_reference_difference = max(max_reference_difference, float(np.max(np.abs(stored - replay))))
            assert reference_result is not None
            diff = float(np.max(np.abs(values - np.asarray([reference_result[c] for c in metrics]))))
            pressure_error = abs(float(result['Ppeak_cmH2O']) -
                                 (float(reference_result['Ppeak_cmH2O']) + peep - 5.0))
            max_invariance_difference = max(max_invariance_difference, diff)
            max_pressure_offset_error = max(max_pressure_offset_error, pressure_error)
            if diff > ATOL or pressure_error > ATOL:
                raise ValueError('The supplied model does not satisfy the pressure-offset invariance check.')
            record = {'reference_scenario_id': int(sid + 1), 'phenotype': str(row.phenotype),
                      **{c: row[c] for c in SETTINGS}, 'VT_ml': float(row.VT_ml),
                      'PBW_kg': float(row.PBW_kg), 'PEEP_cmH2O': peep, 'dt_s': DT,
                      **{c: result[c] for c in [*metrics, 'Ppeak_cmH2O', 'minute_ventilation_L_min']},
                      'reference_Ppeak_cmH2O': reference_result['Ppeak_cmH2O'],
                      'max_matched_metric_abs_difference': diff,
                      'pressure_offset_error_cmH2O': pressure_error}
            record['feasible'] = bool(record['Ppeak_cmH2O'] <= 35.0 and
                                      6.5 <= record['minute_ventilation_L_min'] <= 7.5)
            records.append(record)
    return pd.DataFrame(records), {
        'direct_model_evaluations': len(records), 'reference_scenarios': len(x),
        'maximum_reference_replay_abs_difference': max_reference_difference,
        'maximum_matched_energetic_or_volume_abs_difference': max_invariance_difference,
        'maximum_pressure_offset_error_cmH2O': max_pressure_offset_error,
        'reference_csv_reproduced': True, 'pressure_offset_invariance_verified': True}


def summarize(y: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    summaries, selected = [], []
    baseline_ids: dict[tuple[str, str], set[int]] = {}
    for peep in PEEPS:
        for phenotype in ORDER:
            group = y[y.PEEP_cmH2O.eq(peep) & y.phenotype.eq(phenotype)]
            feasible = group[group.feasible].copy()
            if feasible.empty:
                raise ValueError(f'No feasible settings for {phenotype} at PEEP {peep}.')
            score = np.zeros(len(feasible))
            bounds: dict[str, float] = {}
            for col in ('MP_J_min', 'EII', 'DCES'):
                lo, hi = float(feasible[col].min()), float(feasible[col].max())
                bounds[col + '_normalization_min'] = lo
                bounds[col + '_normalization_max'] = hi
                norm = np.zeros(len(feasible)) if np.isclose(hi, lo) else (feasible[col].to_numpy() - lo) / (hi - lo)
                feasible[col + '_norm'] = norm
                score += norm
            feasible['balanced_score'] = score
            info: dict[str, Any] = {'PEEP_cmH2O': peep, 'phenotype': phenotype,
                'evaluated_scenarios': len(group), 'feasible_scenarios': len(feasible),
                'excluded_by_peak_pressure': len(group) - len(feasible),
                'Ppeak_all_min_cmH2O': float(group.Ppeak_cmH2O.min()),
                'Ppeak_all_max_cmH2O': float(group.Ppeak_cmH2O.max()),
                'Ppeak_feasible_max_cmH2O': float(feasible.Ppeak_cmH2O.max()), **bounds}
            for strategy, col in (('minimum_MP', 'MP_J_min'), ('balanced', 'balanced_score')):
                minimum = float(feasible[col].min())
                tied = feasible[np.isclose(feasible[col], minimum, rtol=RTOL, atol=ATOL)].copy()
                ids = set(tied.reference_scenario_id.astype(int))
                if peep == 5.0:
                    baseline_ids[(phenotype, strategy)] = ids
                unchanged = ids == baseline_ids[(phenotype, strategy)]
                tied['selection'] = strategy
                tied['n_tied_minimizers'] = len(tied)
                tied['settings_unchanged_from_PEEP5'] = unchanged
                selected.append(tied)
                info[strategy + '_n_minimizers'] = len(tied)
                info[strategy + '_settings_unchanged_from_PEEP5'] = unchanged
            summaries.append(info)
    return pd.DataFrame(summaries), pd.concat(selected, ignore_index=True)


def export_docx(summary: pd.DataFrame, selected: pd.DataFrame, out: Path) -> Path:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    import re

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(1.35)
    section.left_margin = section.right_margin = Cm(1.5)
    section.footer_distance = Cm(.6)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(10)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.04

    notation = {'Ppeak': ('P', 'peak'), 'Paw': ('P', 'aw'), 'Ti': ('T', 'i'),
                'VT': ('V', 'T'), 'Eglobal': ('E', 'global'), 'Pabove': ('P', 'above')}
    pattern = re.compile(r'\b(' + '|'.join(notation) + r')\b')
    def text_runs(p, text, *, bold=False, size=10):
        pos = 0
        for match in pattern.finditer(text):
            r = p.add_run(text[pos:match.start()]); r.bold = bold; r.font.size = Pt(size)
            base, sub = notation[match.group()]
            r = p.add_run(base); r.bold = bold; r.font.size = Pt(size)
            r = p.add_run(sub); r.bold = bold; r.font.size = Pt(size); r.font.subscript = True
            pos = match.end()
        r = p.add_run(text[pos:]); r.bold = bold; r.font.size = Pt(size)

    def para(text, bold=False, size=10):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = bold
        text_runs(p, text, bold=bold, size=size)
        return p

    def table(headers, rows, widths, size=9.5):
        t = doc.add_table(rows=1, cols=len(headers))
        t.autofit = False
        t.style = 'Table Grid'
        for col, width in zip(t.columns, widths):
            col.width = Cm(width)
        for values in rows:
            t.add_row()
        for rowid, (row, values) in enumerate(zip(t.rows, [headers, *rows])):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if rowid == 0:
                row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for ci, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
                cell.width = Cm(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
                text_runs(p, str(value), bold=(rowid == 0), size=size)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return t

    para('Supplementary Table S6. Prescribed-PEEP offsets and feasibility in the fixed-mechanics model.', True, 12)
    para('This structural check uses the original reference inspiration-only grid (324 scenarios; 81 per configuration), '
         'with volume reset and a 0.002-s time step. It is not physiological PEEP optimization. '
         'All original simulation outputs and previously reported tables remain unchanged.', size=9.7)
    para('A. Feasibility after direct re-evaluation at four prescribed PEEP values', True, 10.5)
    rows = []
    for peep in PEEPS:
        sub = summary[summary.PEEP_cmH2O.eq(peep)].set_index('phenotype')
        rows.append([f'{peep:g}', *[int(sub.loc[p, 'feasible_scenarios']) for p in ORDER],
                     int(sub.feasible_scenarios.sum()), int(sub.excluded_by_peak_pressure.sum()),
                     f'{sub.Ppeak_all_min_cmH2O.min():.2f}–{sub.Ppeak_all_max_cmH2O.max():.2f}'])
    table(['PEEP\n(cmH₂O)', 'Compliance-\ndominant', 'Resistance-\ndominant', 'Mixed', 'Severe\nmixed',
           'Total\nfeasible (n)', 'Excluded\n(n)', 'Ppeak range before\nfiltering (cmH₂O)'], rows,
          [2.1, 3.7, 3.7, 2.5, 2.5, 3.4, 2.5, 6.3])
    para('Every pressure-offset test evaluated the same 324 VT–RR–Ti–waveform–pause combinations. '
         'Minute ventilation remained 6.5–7.5 L/min. Exclusions reflect only the absolute Ppeak ≤35 cmH₂O criterion. '
         'The displayed pressure ranges include all candidates, including those excluded at higher PEEP.', size=9.5)
    para('B. Structural invariance at matched ventilator settings', True, 10.5)
    para('With resistance, compliance, imposed flow and relative initial volumes fixed, the reference equations give:')
    para('Pabove(t) = Paw(t) − PEEP = [Q(t) + ΣᵢVᵢ(t)/(CᵢRᵢ)] / [Σᵢ1/Rᵢ].', size=10.5)
    para('PEEP is absent from the right-hand side. Replacing 5 cmH₂O by PEEP* therefore gives '
         'Paw*(t) = Paw,5(t) + (PEEP* − 5), with no change in relative compartment volumes, flows, '
         'or energy metrics evaluated above prescribed PEEP. In particular, Ppeak* = Ppeak,5 + PEEP* − 5.')
    para('All 1,296 direct model evaluations reproduced this dependence: at matched settings, global inspiratory '
         'energy, MP, each compartmental positive-inflow energy, each regional energy fraction, EII and DCES '
         'were unchanged. PEEP did not enter the matched-setting objectives; it entered only the absolute '
         'pressure-based feasibility filter.')
    para('This invariance follows from the prescribed model structure and reference-pressure convention. It '
         'does not establish that PEEP has no effect on lung loading in patients. Baseline PEEP-related work '
         'was not added to the study’s above-PEEP inspiratory metric.', size=9.5)

    doc.add_page_break()
    para('Supplementary Table S6 (continued).', True, 12)
    para('C. Minimum-MP and balanced settings after refiltering', True, 10.5)
    para('The listed settings remained selected at PEEP 5, 10, 15 and 20 cmH₂O. Each has VT 4 mL/kg PBW '
         '(280 mL), RR 25 breaths/min and MV 7.0 L/min. MP, EII and DCES did not change with the PEEP offset. '
         'The last columns show the corresponding absolute peak airway pressures.', size=9.7)
    rows = []
    for ph in ORDER:
        for strategy, label in (('minimum_MP', 'Minimum MP'), ('balanced', 'Balanced')):
            subset = selected[selected.phenotype.eq(ph) & selected.selection.eq(strategy)]
            at5 = subset[subset.PEEP_cmH2O.eq(5.)].iloc[0]
            at20 = subset[subset.PEEP_cmH2O.eq(20.)].iloc[0]
            rows.append([LABELS[ph], label, f'{at5.Ti_s:g}', str(at5.waveform).capitalize(),
                         f'{at5.pause_fraction:.1f}', f'{at5.MP_J_min:.3f}', f'{at5.EII:.3f}', f'{at5.DCES:.3f}',
                         f'{at5.Ppeak_cmH2O:.2f}', f'{at20.Ppeak_cmH2O:.2f}'])
    table(['Configuration', 'Selection', 'Ti\n(s)', 'Waveform', 'Pause\nfraction', 'MP\n(J/min)', 'EII', 'DCES',
           'Ppeak\nat PEEP 5', 'Ppeak\nat PEEP 20'], rows,
          [4.1, 2.7, 1.2, 3.3, 1.8, 2.35, 1.8, 1.8, 3.2, 3.95], size=9.3)
    para('Peak airway pressure is expressed in cmH₂O. At PEEP 10 and 15 cmH₂O, add 5 and 10 cmH₂O, '
         'respectively, to the Ppeak values at PEEP 5. All reported selections remained below the pressure limit.', size=9.5)
    para('Selection and numerical checks', True, 10.5)
    para('The original model was directly re-evaluated at each PEEP value. All 324 results at 5 cmH₂O were '
         'checked against the stored primary CSV. The pressure-offset relationship and invariance of '
         'above-PEEP metrics were checked at every candidate setting. Source files were not modified.')
    para('Minimum-MP and balanced strategies were selected separately within each configuration and PEEP value. '
         'For balanced selection, MP, EII and DCES were min–max normalized again within each retained subset; '
         'the original equal-weight sum was unchanged. Ties used relative tolerance 10⁻¹⁰ and absolute tolerance '
         '10⁻¹². All minimum-MP and balanced selections were unique within these tolerances. '
         'Although feasible-set membership and some normalization bounds changed, the selected settings did not. '
         'No assertion is made about invariance of the whole Pareto front.')
    para('Interpretation and limitations', True, 10.5)
    para('PEEP 5, 10, 15 and 20 cmH₂O were used as illustrative offsets, not recommended clinical settings. '
         'The check does not represent PEEP-dependent recruitment, closure, overdistension, changing Rᵢ or Cᵢ, '
         'or patient-specific absolute lung volumes. It was not combined with the passive-expiration or refined-grid '
         'analyses and cannot identify a physiologically preferred PEEP.')
    para('Abbreviations: PEEP, positive end-expiratory pressure; Paw, airway-opening pressure; Pabove, airway pressure '
         'above prescribed PEEP; Ppeak, peak airway-opening pressure; VT, tidal volume; Ti, total inspiratory time; '
         'RR, respiratory rate; MV, minute ventilation; PBW, predicted body weight; MP, global inspiratory mechanical '
         'power above prescribed PEEP; EII, Energy Inequality Index; DCES, Dominant Compartment Energy Share. '
         'EII and DCES are dimensionless. Rᵢ, Cᵢ and Vᵢ denote compartmental resistance, compliance and relative volume.', size=9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Supplementary Table S6 | ').font.size = Pt(8)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    doc.core_properties.title = 'Supplementary Table S6: Prescribed-PEEP offset and feasibility check'
    doc.core_properties.author = 'Tomasz Urbankowski'
    path = out / 'Supplementary_Table_S6_PEEP_offset.docx'
    doc.save(path)
    return path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--input-csv', type=Path, default=Path('MVtarget_primary_feasible_simulation_results.csv'))
    parser.add_argument('--model-script', type=Path, default=Path('multiobjective_ventilation_optimization_scaled_CR_MVtarget.py'))
    parser.add_argument('--output-dir', type=Path, default=Path('peep_offset_sensitivity'))
    parser.add_argument('--no-docx', action='store_true')
    args = parser.parse_args()
    try:
        source, model_path, out = args.input_csv.resolve(), args.model_script.resolve(), args.output_dir.resolve()
        for path in (source, model_path):
            if not path.is_file():
                raise ValueError(f'File not found: {path}')
        out.mkdir(parents=True, exist_ok=True)
        hashes = {path: sha256(path) for path in (source, model_path)}
        x = pd.read_csv(source).reset_index(drop=True)
        validate_input(x)
        y, checks = directly_replay(x, load_model(model_path))
        summary, selected = summarize(y)
        # Baseline selection must agree with the original primary CSV.
        for ph in ORDER:
            sub = x[x.phenotype.eq(ph)]
            for label, col in (('minimum_MP', 'MP_J_min'), ('balanced', 'balanced_score')):
                expected = set(sub[np.isclose(sub[col], sub[col].min(), rtol=RTOL, atol=ATOL)].index + 1)
                actual = set(selected[selected.PEEP_cmH2O.eq(5.) & selected.phenotype.eq(ph) &
                                      selected.selection.eq(label)].reference_scenario_id)
                if expected != actual:
                    raise ValueError(f'Baseline selection mismatch for {ph}/{label}.')
        paths = [out / name for name in ('peep_offset_scenarios.csv', 'peep_offset_summary.csv',
                                        'peep_offset_selected_strategies.csv')]
        if source in paths or model_path in paths:
            raise ValueError('Output path conflicts with a source file.')
        for frame, path in zip((y, summary, selected), paths):
            frame.to_csv(path, index=False)
        if not args.no_docx:
            if not (summary.minimum_MP_n_minimizers.eq(1).all() and summary.balanced_n_minimizers.eq(1).all()
                    and summary.minimum_MP_settings_unchanged_from_PEEP5.all()
                    and summary.balanced_settings_unchanged_from_PEEP5.all()):
                raise ValueError('Selection patterns differ; review/revise the study-specific Table S6 template.')
            paths.append(export_docx(summary, selected, out))
        if any(sha256(path) != before for path, before in hashes.items()):
            raise RuntimeError('A source file was unexpectedly modified.')
        checks.update({'input_name': source.name, 'input_sha256': hashes[source],
                       'model_name': model_path.name, 'model_sha256': hashes[model_path],
                       'PEEP_values_cmH2O': list(PEEPS), 'dt_s': DT, 'rtol': RTOL, 'atol': ATOL,
                       'all_baseline_minimum_MP_and_balanced_selections_reproduced': True,
                       'all_reported_selections_unchanged': bool(summary.minimum_MP_settings_unchanged_from_PEEP5.all()
                                                                 and summary.balanced_settings_unchanged_from_PEEP5.all()),
                       'numpy_version': np.__version__, 'pandas_version': pd.__version__,
                       'physiological_PEEP_optimization_performed': False,
                       'source_files_unchanged': True})
        (out / 'peep_offset_verification.json').write_text(json.dumps(checks, indent=2), encoding='utf-8')
        print(summary[['PEEP_cmH2O', 'phenotype', 'feasible_scenarios', 'excluded_by_peak_pressure']].to_string(index=False))
        print(json.dumps(checks, indent=2))
        for path in paths:
            print(path)
    except (ValueError, OSError, ImportError, RuntimeError) as exc:
        parser.exit(1, f'Error: {exc}\n')


if __name__ == '__main__':
    main()
