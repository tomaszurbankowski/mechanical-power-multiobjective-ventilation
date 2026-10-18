#!/usr/bin/env python3
"""Reproduce Supplementary Table S5: temporal and parameter-grid sensitivity.

This additional analysis uses the ORIGINAL LINEAR INSPIRATION-ONLY, VOLUME-RESET
model. It does not replace previously reported numerical results, reanalyse the
passive-expiration extension, or certify a continuous global optimum/Pareto front.

Run beside the original model and its primary feasible CSV:
    python assess_grid_time_sensitivity.py --output-dir .
Omit the optional Word table:
    python assess_grid_time_sensitivity.py --output-dir . --no-docx

Dependencies: numpy, pandas; python-docx only for the optional Word output.

The original model supplies mechanics and discrete flow waveforms. A batched
NumPy implementation uses the identical Euler/left-sum update order. Exact linear
amplitude scaling avoids redundant integrations at different VT/RR; it is NOT
interpolation of outputs. Every original feasible output is checked against its
stored CSV, and 32 deterministic off-grid cases are checked by direct calls to
the original simulate_breath function. Checks fail rather than silently adapting
to a changed model, dataset, or expected design.
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import itertools
import json
import platform
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

MODEL_NAME = 'multiobjective_ventilation_optimization_scaled_CR_MVtarget.py'
CSV_NAME = 'MVtarget_primary_feasible_simulation_results.csv'
ORDER = ('compliance_dominant', 'resistance_dominant', 'mixed', 'severe_mixed')
LABEL = {'compliance_dominant': 'Compliance-dominant',
         'resistance_dominant': 'Resistance-dominant', 'mixed': 'Mixed',
         'severe_mixed': 'Severe mixed'}
WAVES = ('square', 'decelerating', 'sinusoidal')
OBJECTIVES = ('MP_J_min', 'EII', 'DCES')
METRICS = (*OBJECTIVES, 'Ppeak_cmH2O')
KEYS = ('phenotype', 'VT_ml_per_kg', 'RR_min', 'Ti_s', 'waveform', 'pause_fraction')
RTOL, ATOL = 1e-10, 1e-12
GRID_DT = .0005


def load_model(path: Path) -> Any:
    spec = importlib.util.spec_from_file_location('reference_inspiration_model', path)
    if spec is None or spec.loader is None:
        raise ImportError(f'Cannot load model: {path}')
    model = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(model)
    for name in ('get_phenotypes', 'generate_flow_waveform', 'simulate_breath', 'CMH2O_L_TO_J'):
        if not hasattr(model, name):
            raise ValueError(f'Model lacks {name}.')
    return model


def nested_grids() -> list[tuple[np.ndarray, ...]]:
    axes = [np.array(a, float) for a in
            ([4, 5, 6, 7, 8], [10, 15, 20, 25, 30], [.6, 1., 1.5], [0, .1, .2])]
    result = []
    for _ in range(4):
        result.append(tuple(axes))
        axes = [np.sort(np.r_[a, (a[1:] + a[:-1]) / 2]) for a in axes]
    return result


class Evaluator:
    """Identical discrete reference equations, batched across mechanics/waveforms."""
    def __init__(self, model: Any):
        self.model = model
        self.props = model.get_phenotypes()
        if tuple(self.props) != ORDER:
            raise ValueError('Unexpected mechanical configurations.')
        self.cache: dict[tuple, dict[str, float]] = {}

    def shapes(self, ti: float, pauses: np.ndarray, dt: float) -> None:
        keys = [(ph, float(ti), w, float(pa), float(dt))
                for ph, w, pa in itertools.product(ORDER, WAVES, pauses)]
        keys = [key for key in keys if key not in self.cache]
        if not keys:
            return
        C = np.array([self.props[k[0]]['C'] for k in keys])
        R = np.array([self.props[k[0]]['R'] for k in keys])
        if np.any(C <= 0) or np.any(R <= 0) or C.shape[1] != 6:
            raise ValueError('Expected six positive linear RC compartments.')
        flows = np.array([self.model.generate_flow_waveform(.28, k[1], k[3], k[2], dt=dt)
                          for k in keys]).T
        if not np.allclose(flows.sum(axis=0)*dt, .28, rtol=0, atol=ATOL):
            raise ValueError('Flow waveforms fail the tidal-volume check.')
        V = np.zeros_like(C)
        EC = np.zeros_like(C)
        EG = np.zeros(len(keys))
        peak = np.zeros(len(keys))
        inv_R, inv_CR = 1/R, 1/(C*R)
        denom = inv_R.sum(axis=1)
        for qtot in flows:
            pressure = (qtot + (V*inv_CR).sum(axis=1)) / denom
            q = pressure[:, None]*inv_R - V*inv_CR
            EG += pressure*qtot*dt
            EC += pressure[:, None]*np.maximum(q, 0.)*dt
            V += q*dt
            peak = np.maximum(peak, pressure)
        REF = EC/EC.sum(axis=1)[:, None]
        conversion = float(self.model.CMH2O_L_TO_J)
        for j, key in enumerate(keys):
            self.cache[key] = {
                'MP_J_min': float(EG[j]*conversion*25.),
                'EII': float(np.std(REF[j])/np.mean(REF[j])),
                'DCES': float(REF[j].max()),
                'Ppeak_cmH2O': float(5 + peak[j]),
            }

    def grid(self, axes: tuple[np.ndarray, ...], dt: float) -> pd.DataFrame:
        vt, rr, ti, pause = axes
        # The MV filter depends only on prescribed settings, not simulated outputs.
        pairs = np.array([(a, b) for a in vt for b in rr
                          if 6.5-ATOL <= .07*a*b <= 7.5+ATOL])
        if not len(pairs):
            raise ValueError('No pairs satisfy the MV constraint.')
        for t in ti:
            self.shapes(float(t), pause, dt)
        results = []
        for ph in ORDER:
            rows = []
            for t, w, p in itertools.product(ti, WAVES, pause):
                rows.append({'phenotype': ph, 'Ti_s': float(t), 'waveform': w,
                             'pause_fraction': float(p),
                             **self.cache[(ph, float(t), w, float(p), float(dt))]})
            base = pd.DataFrame(rows)
            n = len(base)
            data = base.iloc[np.repeat(np.arange(n), len(pairs))].reset_index(drop=True)
            data['VT_ml_per_kg'] = np.tile(pairs[:, 0], n)
            data['RR_min'] = np.tile(pairs[:, 1], n)
            scale = data.VT_ml_per_kg/4.
            data['MP_J_min'] *= scale**2 * data.RR_min/25.
            data['Ppeak_cmH2O'] = 5 + (data.Ppeak_cmH2O-5)*scale
            data['minute_ventilation_L_min'] = .07*data.VT_ml_per_kg*data.RR_min
            data = data[data.Ppeak_cmH2O <= 35.].copy()
            results.append(data)
        return pd.concat(results, ignore_index=True)


def summarize(data: pd.DataFrame, run: str, dt: float,
              fixed_bounds: dict | None = None) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    summary, selections, bounds = [], [], {}
    for ph in ORDER:
        sub = data[data.phenotype.eq(ph)].copy()
        low, high = sub[list(OBJECTIVES)].min(), sub[list(OBJECTIVES)].max()
        if not (high > low).all():
            raise ValueError('Degenerate normalization range.')
        bounds[ph] = (low, high)
        sub['balanced_score'] = ((sub[list(OBJECTIVES)]-low)/(high-low)).sum(axis=1)
        if fixed_bounds is not None:
            a, b = fixed_bounds[ph]
            # Deliberately not clipped: this is the original fixed scalar objective.
            sub['fixed_balanced_score'] = ((sub[list(OBJECTIVES)]-a)/(b-a)).sum(axis=1)
        rec = {'run': run, 'dt_s': dt, 'phenotype': ph, 'feasible_n': len(sub)}
        for metric in OBJECTIVES:
            rec[metric+'_minimum'] = float(low[metric])
            rec[metric+'_maximum'] = float(high[metric])
        rec['Ppeak_max_cmH2O'] = float(sub.Ppeak_cmH2O.max())
        local = {}
        for name, objective in [('minimum_MP', 'MP_J_min'), ('balanced', 'balanced_score'),
                                ('minimum_EII', 'EII'), ('minimum_DCES', 'DCES'),
                                ('fixed_reference_balanced', 'fixed_balanced_score')]:
            if objective not in sub:
                continue
            tied = sub[np.isclose(sub[objective], sub[objective].min(), rtol=RTOL, atol=ATOL)]
            # All tie counts are retained; the displayed representative minimizes MP.
            rep = tied.sort_values(['MP_J_min', 'VT_ml_per_kg', 'RR_min',
                                   'Ti_s', 'waveform', 'pause_fraction']).iloc[0]
            record = {'run': run, 'dt_s': dt, 'phenotype': ph, 'strategy': name,
                      'objective_value': float(rep[objective]), 'n_minimizers': len(tied)}
            for col in (*KEYS[1:], 'minute_ventilation_L_min', *METRICS):
                record[col] = rep[col]
            for col in ('VT_ml_per_kg', 'RR_min', 'Ti_s', 'pause_fraction'):
                record[col+'_tied_values'] = ';'.join(f'{v:g}' for v in sorted(tied[col].unique()))
            selections.append(record)
            local[name] = rep
        for metric in OBJECTIVES:
            rec['balanced_vs_minMP_'+metric+'_pct'] = float(
                100*(local['balanced'][metric]/local['minimum_MP'][metric]-1))
        summary.append(rec)
    return pd.DataFrame(summary), pd.DataFrame(selections), bounds


def time_errors(data: pd.DataFrame, reference: pd.DataFrame, run: str,
                dt: float, ref_dt: float) -> list[dict]:
    joined = data.merge(reference, on=list(KEYS), suffixes=('_test', '_ref'), validate='one_to_one')
    if len(joined) != len(data) or len(data) != len(reference):
        raise ValueError('Time-step datasets contain different feasible settings.')
    out = []
    for ph in ORDER:
        sub = joined[joined.phenotype.eq(ph)]
        for metric in METRICS:
            error = abs(sub[metric+'_test']-sub[metric+'_ref'])
            relative = 100*error/abs(sub[metric+'_ref'])
            out.append({'run': run, 'dt_s': dt, 'reference_dt_s': ref_dt,
                        'phenotype': ph, 'n': len(sub), 'metric': metric,
                        'maximum_absolute_difference': float(error.max()),
                        'maximum_relative_difference_pct': float(relative.max())})
    return out


def export_table(out: Path, definitions: pd.DataFrame, grid_summary: pd.DataFrame,
                 grid_selections: pd.DataFrame, temporal: pd.DataFrame,
                 boundary: pd.DataFrame) -> Path:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    sec.top_margin = sec.bottom_margin = Cm(1.35)
    sec.left_margin = sec.right_margin = Cm(1.35)
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(9.5)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.line_spacing = 1.05

    def para(text: str, bold: bool = False, size: float = 9.5):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = bold
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        return p

    def table(headers: list[str], rows: list[list[str]], widths: list[float], fontsize: float=9):
        t = doc.add_table(rows=1, cols=len(headers))
        t.autofit = False
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = 'Table Grid'
        for c,w in zip(t.columns,widths): c.width=Cm(w)
        for cell,text,w in zip(t.rows[0].cells,headers,widths):
            cell.width=Cm(w);cell.text=text
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for c,text,w in zip(t.add_row().cells,row,widths):c.width=Cm(w);c.text=text
        for i,row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for j,c in enumerate(row.cells):
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcPr=c._tc.get_or_add_tcPr()
                margins=OxmlElement('w:tcMar')
                for tag,val in [('top','65'),('bottom','65'),('left','80'),('right','80')]:
                    el=OxmlElement('w:'+tag);el.set(qn('w:w'),val);el.set(qn('w:type'),'dxa');margins.append(el)
                tcPr.append(margins)
                if i==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAEFF4');tcPr.append(shade)
                for p in c.paragraphs:
                    p.paragraph_format.space_after=Pt(0)
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:r.font.size=Pt(fontsize);r.bold=i==0
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
        return t

    para('Supplementary Table S5. Sensitivity to temporal and ventilator-setting grid resolution.', True, 12)
    para('All tests below use the reference linear inspiration-only model with volume reset, not the passive-expiration extension. '
         'Previously reported tables and simulation outputs are retained unchanged.')
    para('A. Nested parameter grids', True, 10.5)
    rows=[]
    for _,r in definitions.iterrows():
        rows.append([r['grid'],r['VT_step_mL_kg'],r['RR_step_min'],r['Ti_gaps_s'],r['pause_step'],
                     str(int(r['full_combinations'])),str(int(r['feasible_combinations'])),str(int(r['feasible_per_configuration']))])
    table(['Grid','VT step\n(mL/kg)','RR step\n(breaths/min)','Ti interval\nlengths (s)','Pause-fraction\nstep',
           'Full grid\ncombinations','Feasible\ncombinations','Feasible per\nconfiguration'], rows,
          [1.5,2.7,3.0,3.3,3.1,4.0,3.5,3.2])
    para('G0 contains the original VT, RR, Ti and pause levels. G1-G3 successively insert the midpoint of every interval on each axis. '
         'Ranges remain VT 4-8 mL/kg, RR 10-30 breaths/min, Ti 0.6-1.5 s and pause fraction 0-0.2. '
         'The same three categorical waveforms are used throughout. Ti has unequal original interval lengths (0.4 and 0.5 s).')
    para('All grid comparisons use a common time step of 0.0005 s; G0 is therefore a recomputed control for parameter refinement, '
         'not the original 0.002-s output. PBW is 70 kg and PEEP is 5 cmH₂O. Feasibility is MV 6.5-7.5 L/min and Ppeak ≤35 cmH₂O. '
         'No MV-admissible scenario was removed by the peak-pressure criterion.')
    para('Computation and selection', True, 10.5)
    para('New parameter combinations are evaluated from the original discrete equations, without interpolation of responses. '
         'At fixed mechanics, timing and waveform, the linear zero-reset model scales pressure above PEEP and volume with VT and energy with VT squared. '
         'This exact amplitude relation is used to avoid redundant integrations; MP additionally scales with RR. '
         'The batched implementation was checked against all 324 stored reference scenarios and 32 direct off-grid calls to the original model.')
    para('Minimum-MP and balanced settings are selected within each configuration and grid. Balanced normalization is recomputed within each grid. '
         'A separate check freezes normalization to the G0 bounds, without clipping out-of-range normalized values. '
         'Selection tolerances are relative 1e-10 and absolute 1e-12. The reported minimum-MP and balanced solutions are unique within these tolerances. '
         'These tests do not constitute a grid-convergence test of Pareto-front membership or a continuous global-optimality certificate.')

    doc.add_page_break()
    para('Supplementary Table S5 (continued).', True, 12)
    para('B. Minimum-MP and balanced selections across the nested grids', True, 10.5)
    para('All listed selections use VT 4 mL/kg. G0/G1 select RR 25/min (MV 7.0 L/min); G2/G3 select RR 23.75/min (MV 6.65 L/min). '
         'Minimum-MP always selects Ti 1.5 s, square flow and no pause. Other columns describe the balanced solution using grid-specific normalization.')
    rows=[]
    for ph in ORDER:
        for g in ['G0','G1','G2','G3']:
            sub=grid_selections[grid_selections.phenotype.eq(ph)&grid_selections.run.eq(g)]
            mp=sub[sub.strategy.eq('minimum_MP')].iloc[0];b=sub[sub.strategy.eq('balanced')].iloc[0]
            rows.append([LABEL[ph],g,f'{mp.MP_J_min:.3f}',f'{b.Ti_s:g}',b.waveform.capitalize(),f'{b.pause_fraction:.4f}',
                         f'{b.MP_J_min:.3f}',f'{b.EII:.5f}',f'{b.DCES:.5f}'])
    table(['Configuration','Grid','Minimum MP\n(J/min)','Balanced\nTi (s)','Balanced\nwaveform','Balanced\npause fraction',
           'Balanced MP\n(J/min)','Balanced\nEII','Balanced\nDCES'],rows,[4.0,1.0,3.0,2.0,3.7,3.0,3.0,2.8,2.8],8.7)
    para('Relative to G0 at the same time step, G3 reduces minimum MP by 5.0% in each configuration because lower admissible minute ventilation '
         'becomes accessible. G3 preserves the direction of the balanced-versus-minimum-MP trade-off: lower EII and DCES at higher MP. '
         'Both fixed-G0 and grid-specific normalization select the same G3 settings. At intermediate grid G2, the selected pause fractions differ '
         'between normalization conventions (resistance-dominant: 0.125 versus 0.15; severe mixed: 0.025 versus 0.05).')
    para('Scalar objective minima, all selection records and normalization ranges are reported in the accompanying CSV files. '
         'G2 and G3 have identical minimum-MP settings, but the resistance-dominant and severe-mixed balanced pause fractions still change. '
         'No claim of exact parameter convergence is made.')

    doc.add_page_break()
    para('Supplementary Table S5 (continued).', True, 12)
    para('C. Temporal resolution at matched settings', True, 10.5)
    rows=[]
    for run in ['original_grid','finest_grid']:
        x=temporal[temporal.run.eq(run)]
        for dt in sorted(x.dt_s.unique(),reverse=True):
            z=x[x.dt_s.eq(dt)];refdt=float(z.reference_dt_s.iloc[0])
            if dt==refdt:continue
            vals=[]
            for met in METRICS:
                e=z[z.metric.eq(met)]
                vals.append(f'{e.maximum_relative_difference_pct.max():.4f}%\n({e.maximum_absolute_difference.max():.6f})')
            rows.append(['G0 (324)' if run=='original_grid' else 'G3 (454,308)', f'{dt:g}',f'{refdt:g}',*vals])
    table(['Grid (n)','Test dt\n(s)','Reference dt\n(s)','MP difference\n% (J/min)','EII difference\n% (absolute)',
           'DCES difference\n% (absolute)','Ppeak difference\n% (cmH₂O)'],rows,[3.7,2.2,2.6,4.4,4.3,4.3,4.4],8.6)
    para('Each cell gives the maximum relative difference across all matched settings and, in parentheses, the maximum absolute difference; '
         'the two maxima need not arise from the same scenario. The smallest tested step is a numerical reference, not an exact solution. '
         'Within G0, minimum-MP and balanced settings are unchanged across dt 0.002, 0.001, 0.0005 and 0.00025 s. '
         'Within G3, these selections are unchanged across dt 0.0005, 0.00025 and 0.000125 s. Waveform sampling and integration follow the original implementation, '
         'including rounding active-step counts; temporal differences therefore reflect both discretizations.')
    para('D. Feasible-boundary check: why finite-grid stability is not a continuous-optimality proof', True, 10.5)
    rows=[]
    for _,r in boundary.iterrows():
        rows.append([LABEL[r.phenotype],f'{r.grid_MP_J_min:.3f}',f'{r.boundary_MP_J_min:.3f}',f'{r.relative_to_G0_pct:.4f}%'])
    table(['Configuration','G3 minimum MP\n(J/min)','Same waveform/timing at MV 6.5\nMP (J/min)','Change versus G0\n(%)'],rows,[6.6,6.0,8.0,5.7])
    para('At VT 4 mL/kg, an admissible RR of 6.5/0.28 = 23.2142857 breaths/min reaches the exact MV lower boundary. '
         'Keeping Ti 1.5 s, square flow and no pause, this lowers MP by 7.1429% versus G0 at the same dt of 0.0005 s, '
         'compared with the 5.0% reduction reached by G3. These additional directly evaluated boundary settings demonstrate that even G3 is not '
         'a continuous optimum. They are not a search over all continuous inspiratory waveforms.')
    para('MP, global inspiratory mechanical power above external PEEP; EII, Energy Inequality Index; DCES, Dominant Compartment Energy Share; '
         'VT, tidal volume; RR, respiratory rate; Ti, total inspiratory time; MV, minute ventilation; PBW, predicted body weight; '
         'PEEP, positive end-expiratory pressure; Ppeak, peak airway-opening pressure. EII and DCES are dimensionless.',size=8.5)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Supplementary Table S5 | ').font.size=Pt(8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    # Typeset recurring mathematical subscripts without changing ordinary words.
    import re
    pattern=re.compile(r'\b(Ppeak|Ti|VT)\b')
    mapping={'Ppeak':('P','peak'),'Ti':('T','i'),'VT':('V','T')}
    pars=list(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:pars.extend(cell.paragraphs)
    from copy import deepcopy
    from docx.text.run import Run
    for p in pars:
        for run in list(p.runs):
            text=run.text;matches=list(pattern.finditer(text))
            if not matches:continue
            cursor=0;pieces=[]
            for match in matches:
                pieces.append((text[cursor:match.start()],False));a,b=mapping[match.group()]
                pieces.extend([(a,False),(b,True)]);cursor=match.end()
            pieces.append((text[cursor:],False))
            for text,sub in pieces:
                if not text:continue
                el=OxmlElement('w:r')
                if run._r.rPr is not None:el.append(deepcopy(run._r.rPr))
                run._r.addprevious(el);nr=Run(el,p);nr.text=text;nr.font.subscript=sub
            run._r.getparent().remove(run._r)
    doc.core_properties.author='Tomasz Urbankowski'
    doc.core_properties.title='Supplementary Table S5: Temporal and parameter-grid sensitivity'
    target=out/'Supplementary_Table_S5_numerical_resolution.docx'
    doc.save(target)
    return target


def main() -> None:
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--model-script',type=Path,default=Path(MODEL_NAME))
    parser.add_argument('--input-csv',type=Path,default=Path(CSV_NAME))
    parser.add_argument('--output-dir',type=Path,default=Path('numerical_resolution_sensitivity'))
    parser.add_argument('--no-docx',action='store_true')
    args=parser.parse_args()
    model_path, source=args.model_script.resolve(),args.input_csv.resolve()
    for path in (model_path,source):
        if not path.is_file():parser.error(f'Missing input: {path}')
    input_hash=hashlib.sha256(source.read_bytes()).hexdigest()
    out=args.output_dir.resolve();out.mkdir(parents=True,exist_ok=True)
    model=load_model(model_path); ev=Evaluator(model);grids=nested_grids()
    stored=pd.read_csv(source)
    if len(stored)!=324 or stored.groupby('phenotype').size().to_dict()!=dict.fromkeys(ORDER,81):
        raise ValueError('Expected the original 324 primary feasible scenarios.')
    baseline=ev.grid(grids[0],.002)
    joined=baseline.merge(stored,on=list(KEYS),suffixes=('_new','_stored'),validate='one_to_one')
    if len(joined)!=324:raise ValueError('Original settings do not match stored CSV.')
    baseline_errors={}
    for metric in METRICS:
        a,b=joined[metric+'_new'],joined[metric+'_stored']
        if not np.allclose(a,b,rtol=RTOL,atol=ATOL):raise ValueError(f'Original {metric} not reproduced.')
        baseline_errors[metric]=float(abs(a-b).max())

    defs=[];allsummary=[];allselected=[];fixed=None
    for level,axes in enumerate(grids):
        data=ev.grid(axes,GRID_DT)
        sm,se,bounds=summarize(data,f'G{level}',GRID_DT,fixed)
        if level==0:fixed=bounds
        allsummary.append(sm);allselected.append(se)
        v,r,t,p=axes
        step=lambda a:'; '.join(f'{x:g}' for x in np.unique(np.round(np.diff(a),12)))
        defs.append({'grid':f'G{level}','VT_step_mL_kg':step(v),'RR_step_min':step(r),
                     'Ti_gaps_s':step(t),'pause_step':step(p),
                     'full_combinations':int(4*3*len(v)*len(r)*len(t)*len(p)),
                     'feasible_combinations':len(data),'feasible_per_configuration':len(data)//4,
                     'VT_levels':';'.join(f'{x:g}' for x in v),'RR_levels':';'.join(f'{x:g}' for x in r),
                     'Ti_levels':';'.join(f'{x:g}' for x in t),'pause_levels':';'.join(f'{x:g}' for x in p),
                     'dt_s':GRID_DT})
    definition=pd.DataFrame(defs);summary=pd.concat(allsummary,ignore_index=True)
    selections=pd.concat(allselected,ignore_index=True)
    if definition.feasible_combinations.tolist()!=[324,2700,34992,454308]:
        raise ValueError('Unexpected grid counts; update the documented design before reuse.')
    if not (summary.balanced_vs_minMP_MP_J_min_pct.gt(0)&summary.balanced_vs_minMP_EII_pct.lt(0)&
            summary.balanced_vs_minMP_DCES_pct.lt(0)).all():
        raise ValueError('The stated trade-off direction was not reproduced.')

    temporal_rows=[];temporal_selected=[];direct_errors={c:0. for c in METRICS}
    for run,axes,steps in [('original_grid',grids[0],[.002,.001,.0005,.00025]),
                           ('finest_grid',grids[3],[.0005,.00025,.000125])]:
        refdt=steps[-1];ref=ev.grid(axes,refdt)
        selected_signatures=[]
        for dt in steps:
            data=ref if dt==refdt else ev.grid(axes,dt)
            temporal_rows.extend(time_errors(data,ref,run,dt,refdt))
            _,ss,_=summarize(data,run,dt)
            ss=ss[ss.strategy.isin(['minimum_MP','balanced'])]
            temporal_selected.append(ss)
            selected_signatures.append(ss[['phenotype','strategy',*KEYS[1:]]].reset_index(drop=True))
        for a in selected_signatures[1:]:
            pd.testing.assert_frame_equal(a,selected_signatures[0])
        if run=='finest_grid':
            for index in np.linspace(0,len(ref)-1,32,dtype=int):
                row=ref.iloc[index]
                ans=model.simulate_breath(**ev.props[row.phenotype],vt_l=float(row.VT_ml_per_kg)*.07,
                    rr=float(row.RR_min),ti_s=float(row.Ti_s),waveform=row.waveform,
                    pause_fraction=float(row.pause_fraction),dt=refdt)
                for metric in METRICS:
                    if not np.isclose(row[metric],ans[metric],rtol=RTOL,atol=ATOL):
                        raise ValueError(f'Off-grid direct check failed for {metric}.')
                    direct_errors[metric]=max(direct_errors[metric],float(abs(row[metric]-ans[metric])))
    temporal=pd.DataFrame(temporal_rows)
    tselections=pd.concat(temporal_selected,ignore_index=True)

    boundary=[]
    for ph in ORDER:
        g3=selections[selections.run.eq('G3')&selections.phenotype.eq(ph)&selections.strategy.eq('minimum_MP')].iloc[0]
        g0=selections[selections.run.eq('G0')&selections.phenotype.eq(ph)&selections.strategy.eq('minimum_MP')].iloc[0]
        rr=6.5/.28
        ans=model.simulate_breath(**ev.props[ph],vt_l=.28,rr=rr,ti_s=1.5,waveform='square',pause_fraction=0.,dt=GRID_DT)
        if ans['Ppeak_cmH2O']>35:raise ValueError('Boundary case not feasible.')
        boundary.append({'phenotype':ph,'dt_s':GRID_DT,'VT_ml_per_kg':4,'RR_min':rr,
                         'Ti_s':1.5,'waveform':'square','pause_fraction':0,'MV_L_min':6.5,
                         'grid_MP_J_min':float(g3.MP_J_min),'boundary_MP_J_min':ans['MP_J_min'],
                         'relative_to_G0_pct':100*(ans['MP_J_min']/g0.MP_J_min-1),
                         'EII':ans['EII'],'DCES':ans['DCES'],'Ppeak_cmH2O':ans['Ppeak_cmH2O']})
    boundary=pd.DataFrame(boundary)
    outputs={'numerical_grid_definition.csv':definition,'numerical_grid_summary.csv':summary,
             'numerical_grid_selected_strategies.csv':selections,'numerical_time_step_summary.csv':temporal,
             'numerical_time_step_selected_strategies.csv':tselections,'numerical_boundary_check.csv':boundary}
    for name,df in outputs.items():
        if (out/name).resolve()==source:raise ValueError('Input/output collision.')
        df.to_csv(out/name,index=False)
    if hashlib.sha256(source.read_bytes()).hexdigest()!=input_hash:raise RuntimeError('Input CSV changed.')
    report={'original_csv_sha256':input_hash,'model_sha256':hashlib.sha256(model_path.read_bytes()).hexdigest(),
            'reference_model':'linear inspiration-only with volume reset','original_rows_verified':324,
            'maximum_baseline_absolute_difference':baseline_errors,'direct_off_grid_checks':32,
            'maximum_direct_off_grid_absolute_difference':direct_errors,
            'minimum_MP_and_balanced_selection_stable_over_time_steps':True,
            'original_outputs_modified':False,'continuous_global_optimum_proven':False,
            'pareto_front_convergence_tested':False,'passive_expiration_grid_refinement_tested':False,
            'selection_rtol':RTOL,'selection_atol':ATOL,'python':platform.python_version(),
            'numpy':np.__version__,'pandas':pd.__version__}
    (out/'numerical_resolution_verification.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    if not args.no_docx:
        try:print(export_table(out,definition,summary,selections,temporal,boundary))
        except ImportError as exc:raise SystemExit('Install python-docx or use --no-docx; CSVs were saved.') from exc
    print(definition[['grid','full_combinations','feasible_combinations']].to_string(index=False))
    print('Original dataset reproduced. New results are separate; no continuous optimality is claimed.')

if __name__=='__main__':
    main()
