#!/usr/bin/env python3
"""Compartmental inspiratory work, power, and RC energy-accounting diagnostics.

Reproduce Supplementary Table S7 from the original 2,700-scenario reference
inspiration-only (volume-reset) model, including its primary (324) and broader
minute-ventilation (1,944) feasible subsets. No original files are modified.

Run in the repository root:
    python assess_compartment_energy_power.py --output-dir . --no-docx
Omit --no-docx to also create Supplementary Table S7 (requires python-docx).

Eplus is the manuscript's original positive-inflow work, not net work or tissue
dissipation. MPplus_i = RR * Eplus_i is its rate per minute, NOT an instantaneous
peak or a whole-cycle dissipation rate. Positive-inflow shares, EII and DCES are
unchanged by multiplying all six branch energies by the same RR in a scenario.

Signed work, branch-resistive dissipation, elastic storage above the prescribed-
PEEP equilibrium, and outward work are reported separately. At finite Euler dt:
    Wnet_i = D_i + delta_U_i - euler_increment_i
    Eplus_i - Eminus_i = Wnet_i
The Euler increment is a numerical energy-accounting term, not physical loss.
Only inspiration (including any end-inspiratory pause) is audited here.
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import itertools
import json
from pathlib import Path
import platform
import sys
from types import ModuleType
from typing import Any

import numpy as np
import pandas as pd

DT = 0.002
K = 0.0980665
RTOL, ATOL = 1e-10, 1e-12
ORDER = ('compliance_dominant', 'resistance_dominant', 'mixed', 'severe_mixed')
LABELS = dict(zip(ORDER, ('Compliance-dominant', 'Resistance-dominant', 'Mixed', 'Severe mixed')))
KEYS = ['phenotype', 'VT_ml_per_kg', 'RR_min', 'Ti_s', 'waveform', 'pause_fraction', 'PEEP_cmH2O']
BASE_NAME = 'multiobjective_ventilation_optimization_scaled_CR_MVtarget.py'
CSV_NAME = 'MVtarget_primary_all_simulation_results.csv'
FIELDS = ['global_energy_J', 'MP_J_min', 'EII', 'DCES', 'Ppeak_cmH2O',
          'Pplat_cmH2O', 'driving_pressure_cmH2O', 'minute_ventilation_L_min',
          'end_inspiratory_volume_L'] + [f'{pre}{i}' for i in range(1, 7)
                                        for pre in ('REF_', 'comp_energy_J_')]


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_model(path: Path) -> ModuleType:
    spec = importlib.util.spec_from_file_location('reference_ventilation_model', path)
    if spec is None or spec.loader is None:
        raise ValueError(f'Cannot import model: {path}')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    for name in ('get_phenotypes', 'generate_flow_waveform', 'simulate_breath',
                 'add_normalized_columns'):
        if not callable(getattr(module, name, None)):
            raise ValueError(f'Missing expected reference function: {name}')
    return module


def validate_design(df: pd.DataFrame) -> None:
    required = set(KEYS + FIELDS + ['PBW_kg', 'VT_ml'])
    if missing := required - set(df.columns):
        raise ValueError(f'Missing input columns: {sorted(missing)}')
    if len(df) != 2700:
        raise ValueError('Expected the original all-scenario reference CSV with 2,700 rows.')
    expected = {(p, v, r, t, w, a, 5.0) for p, v, r, t, w, a in itertools.product(
        ORDER, (4, 5, 6, 7, 8), (10, 15, 20, 25, 30), (0.6, 1., 1.5),
        ('square', 'decelerating', 'sinusoidal'), (0., 0.1, 0.2))}
    observed = list(df[KEYS].itertuples(index=False, name=None))
    if len(set(observed)) != len(observed) or set(observed) != expected:
        raise ValueError('Input is not the original 2,700-setting grid.')
    if not np.allclose(df.PBW_kg, 70.0, rtol=0, atol=ATOL):
        raise ValueError('Expected PBW 70 kg.')
    if not np.allclose(df.VT_ml, df.VT_ml_per_kg * df.PBW_kg, rtol=0, atol=ATOL):
        raise ValueError('Inconsistent tidal-volume units.')
    if not np.isfinite(df.select_dtypes(include='number').to_numpy()).all():
        raise ValueError('Non-finite input values.')


def replay_diagnostics(model: ModuleType, C: np.ndarray, R: np.ndarray,
                       vt_l: float, ti_s: float, waveform: str,
                       pause_fraction: float) -> dict[str, Any]:
    """Original Euler dynamics and quadrature, with extra diagnostic accumulators."""
    qtotal = model.generate_flow_waveform(vt_l, ti_s, pause_fraction, waveform, dt=DT)
    v = np.zeros(6)
    eplus, eminus, wnet, dissip, num = (np.zeros(6) for _ in range(5))
    inv_cr, inv_r = 1.0/(C*R), 1.0/R
    denom = np.sum(inv_r)
    global_work = 0.0
    max_pressure = -np.inf
    max_flow_residual = 0.0
    for qt in qtotal:
        p = (qt + np.sum(v * inv_cr)) / denom
        q = p * inv_r - v * inv_cr
        global_work += p * qt * DT
        # Evaluate every energy product BEFORE updating volume (as in the source).
        eplus += p * np.maximum(q, 0.0) * DT
        eminus += p * np.maximum(-q, 0.0) * DT
        wnet += p * q * DT
        dissip += R * q * q * DT
        num += (q * DT)**2 / (2*C)
        v += q * DT
        max_pressure = max(max_pressure, float(p))
        max_flow_residual = max(max_flow_residual, float(abs(np.sum(q)-qt)))
    delta_u = K * v*v/(2*C)  # initial volume is exactly zero
    return dict(Eplus=eplus*K, Eminus=eminus*K, Wnet=wnet*K, D=dissip*K,
                delta_U=delta_u, euler_increment=num*K, V_end=v,
                global_energy_J=global_work*K, max_p_above=max_pressure,
                last_p_above=float(p), max_flow_residual_L_s=max_flow_residual)


def analyze(model: ModuleType, df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, Any]]:
    mechanics = model.get_phenotypes()
    cache: dict[tuple, dict[str, Any]] = {}
    records, components = [], []
    reference_diff = dict.fromkeys(FIELDS, 0.0)
    maxima = dict(flow_conservation_L_s=0., volume_conservation_L=0.,
                  signed_work_sum_J=0., positive_minus_negative_J=0.,
                  discrete_branch_balance_J=0., share_conversion=0.,
                  EII_power_conversion=0., DCES_power_conversion=0.)
    for serial, (_, row) in enumerate(df.iterrows(), start=1):
        C = np.asarray(mechanics[row.phenotype]['C'], dtype=float)
        R = np.asarray(mechanics[row.phenotype]['R'], dtype=float)
        vt_l = float(row.VT_ml_per_kg * row.PBW_kg / 1000.0)
        key = (row.phenotype, row.VT_ml_per_kg, row.Ti_s, row.waveform, row.pause_fraction)
        if key not in cache:
            cache[key] = replay_diagnostics(model, C, R, vt_l, row.Ti_s,
                                            row.waveform, row.pause_fraction)
        d = cache[key]
        eplus = d['Eplus']
        ref = eplus / eplus.sum()
        rr = float(row.RR_min)
        expected = dict(global_energy_J=d['global_energy_J'],
                        MP_J_min=d['global_energy_J']*rr,
                        EII=np.std(ref)/np.mean(ref), DCES=np.max(ref),
                        Ppeak_cmH2O=5.0+d['max_p_above'], Pplat_cmH2O=5.0+d['last_p_above'],
                        driving_pressure_cmH2O=(5.0+d['last_p_above'])-5.0,
                        minute_ventilation_L_min=vt_l*rr, end_inspiratory_volume_L=np.sum(d['V_end']))
        for i in range(6):
            expected[f'REF_{i+1}'] = ref[i]
            expected[f'comp_energy_J_{i+1}'] = eplus[i]
        for field in FIELDS:
            delta = float(abs(float(row[field])-expected[field]))
            reference_diff[field] = max(reference_diff[field], delta)
            if not np.isclose(row[field], expected[field], rtol=RTOL, atol=ATOL):
                raise ValueError(f'Reference mismatch in scenario {serial}, {field}: {row[field]} vs {expected[field]}')
        # Use stored E_i for power (not a new rounded or redefined energy measure).
        stored_e = row[[f'comp_energy_J_{i}' for i in range(1,7)]].to_numpy(dtype=float)
        power = stored_e * rr
        power_shares = power / power.sum()
        out = dict(scenario_id=serial, **row.to_dict())
        out.update(Emax_positive_J_per_breath=float(np.max(stored_e)),
                   MPmax_positive_J_min=float(np.max(power)),
                   Eplus_sum_J_per_breath=float(stored_e.sum()),
                   MPplus_sum_J_min=float(power.sum()),
                   EII_from_power=float(np.std(power_shares)/np.mean(power_shares)),
                   DCES_from_power=float(np.max(power_shares)))
        for name in ('Eminus','Wnet','D','delta_U','euler_increment'):
            out[f'{name}_sum_J_per_breath'] = float(np.sum(d[name]))
        for i in range(6):
            out[f'MPplus_{i+1}_J_min'] = float(power[i])
            components.append(dict(scenario_id=serial, phenotype=row.phenotype,
                compartment=i+1, RR_min=rr, R_cmH2O_s_L=R[i], C_L_cmH2O=C[i],
                V_end_L=d['V_end'][i], Eplus_J_per_breath=stored_e[i],
                Eminus_J_per_breath=d['Eminus'][i], Wnet_J_per_breath=d['Wnet'][i],
                D_inspiratory_J_per_breath=d['D'][i],
                delta_U_above_PEEP_J_per_breath=d['delta_U'][i],
                euler_increment_J_per_breath=d['euler_increment'][i],
                MPplus_J_min=power[i], Wnet_rate_J_min=rr*d['Wnet'][i],
                D_inspiratory_rate_J_min=rr*d['D'][i],
                delta_U_rate_J_min=rr*d['delta_U'][i],
                REF=float(stored_e[i]/stored_e.sum())))
        maxima['flow_conservation_L_s'] = max(maxima['flow_conservation_L_s'],d['max_flow_residual_L_s'])
        maxima['volume_conservation_L'] = max(maxima['volume_conservation_L'],float(abs(np.sum(d['V_end'])-vt_l)))
        maxima['signed_work_sum_J'] = max(maxima['signed_work_sum_J'],float(abs(np.sum(d['Wnet'])-d['global_energy_J'])))
        maxima['positive_minus_negative_J'] = max(maxima['positive_minus_negative_J'],float(np.max(abs(d['Eplus']-d['Eminus']-d['Wnet']))))
        maxima['discrete_branch_balance_J'] = max(maxima['discrete_branch_balance_J'],float(np.max(abs(d['Wnet']-d['D']-d['delta_U']+d['euler_increment']))))
        maxima['share_conversion'] = max(maxima['share_conversion'],float(np.max(abs(stored_e/stored_e.sum()-power_shares))))
        maxima['EII_power_conversion'] = max(maxima['EII_power_conversion'],float(abs(out['EII_from_power']-row.EII)))
        maxima['DCES_power_conversion'] = max(maxima['DCES_power_conversion'],float(abs(out['DCES_from_power']-row.DCES)))
        records.append(out)
    if any(v > 1e-10 for v in maxima.values()):
        raise ValueError(f'Diagnostic identity failed: {maxima}')
    return pd.DataFrame(records), pd.DataFrame(components), dict(
        scenarios_replayed=2700, unique_inspirations_replayed=len(cache),
        max_absolute_original_output_differences=reference_diff,
        max_absolute_identity_residuals=maxima)


def summarize(model: ModuleType, all_df: pd.DataFrame, source_dir: Path) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, dict]:
    selected, summaries, comparisons = [], [], []
    controls = {}
    for analysis_name, lo, hi, count, stored_name in (
        ('primary',6.5,7.5,324,'MVtarget_primary_feasible_simulation_results.csv'),
        ('broader_MV',4.,12.,1944,'MVtarget_sensitivity_feasible_simulation_results.csv')):
        mask = (all_df.Ppeak_cmH2O<=35) & all_df.minute_ventilation_L_min.between(lo,hi)
        all_df[f'feasible_{analysis_name}'] = mask
        sub = all_df.loc[mask].copy()
        if len(sub)!=count:
            raise ValueError(f'Unexpected feasible count for {analysis_name}: {len(sub)}')
        sub = model.add_normalized_columns(sub, ['MP_J_min','EII','DCES'])
        sub['balanced_score'] = sub.MP_J_min_norm+sub.EII_norm+sub.DCES_norm
        stored_path = source_dir / stored_name
        if stored_path.is_file():
            old = pd.read_csv(stored_path, float_precision='round_trip')
            m = sub.merge(old,on=KEYS,suffixes=('_new','_old'),validate='one_to_one')
            if len(m)!=count or len(old)!=count:
                raise ValueError(f'Feasible membership mismatch: {analysis_name}')
            for f in FIELDS+['balanced_score']:
                if not np.allclose(m[f'{f}_new'],m[f'{f}_old'],rtol=RTOL,atol=ATOL):
                    raise ValueError(f'Feasible set mismatch: {analysis_name}, {f}')
            controls[analysis_name] = f'Checked all {count} rows, original metrics, and balanced scores against stored feasible CSV.'
        else:
            controls[analysis_name] = 'Reconstructed from all-scenario CSV; no separate feasible CSV available for comparison.'
        for p in ORDER:
            group = sub.loc[sub.phenotype==p]
            objective_map = {'minimum_MP':'MP_J_min','balanced':'balanced_score',
                             'minimum_maximum_energy':'Emax_positive_J_per_breath',
                             'minimum_maximum_power':'MPmax_positive_J_min'}
            chosen = {}
            for strategy, objective in objective_map.items():
                vmin = group[objective].min()
                ties = group.loc[np.isclose(group[objective],vmin,rtol=RTOL,atol=ATOL)].copy()
                chosen[strategy]=ties
                for tie_n,(_,row) in enumerate(ties.iterrows(),start=1):
                    selected.append(dict(analysis_set=analysis_name,strategy=strategy,
                                         tied_minima_count=len(ties),tie_number=tie_n,
                                         minimized_objective=objective,minimum_value=vmin,**row.to_dict()))
            amin,b = chosen['minimum_MP'].iloc[0],chosen['balanced'].iloc[0]
            cmp = dict(analysis_set=analysis_name,phenotype=p,
                       minimum_MP_scenario=int(amin.scenario_id),balanced_scenario=int(b.scenario_id))
            for f in ('MP_J_min','EII','DCES','Emax_positive_J_per_breath','MPmax_positive_J_min',
                      'D_sum_J_per_breath','delta_U_sum_J_per_breath'):
                cmp[f'{f}_minimum_MP']=amin[f]; cmp[f'{f}_balanced']=b[f]
                cmp[f'{f}_change_percent']=100*(b[f]/amin[f]-1)
            comparisons.append(cmp)
            summ = dict(analysis_set=analysis_name,phenotype=p,n=len(group),
                        minimum_maximum_energy_ties=len(chosen['minimum_maximum_energy']),
                        energy_minimizer_RR_values=';'.join(str(float(x)) for x in sorted(chosen['minimum_maximum_energy'].RR_min.unique())),
                        minimum_maximum_power_ties=len(chosen['minimum_maximum_power']),
                        energy_minimizer_power_min=chosen['minimum_maximum_energy'].MPmax_positive_J_min.min(),
                        energy_minimizer_power_max=chosen['minimum_maximum_energy'].MPmax_positive_J_min.max(),
                        max_euler_increment_relative_to_global_work_percent=100*(group.euler_increment_sum_J_per_breath/group.global_energy_J).max())
            for f in ('Emax_positive_J_per_breath','MPmax_positive_J_min','Eminus_sum_J_per_breath',
                      'D_sum_J_per_breath','delta_U_sum_J_per_breath','euler_increment_sum_J_per_breath'):
                summ[f'{f}_min']=group[f].min(); summ[f'{f}_max']=group[f].max()
            summaries.append(summ)
    return pd.DataFrame(selected),pd.DataFrame(summaries),pd.DataFrame(comparisons),controls


def main() -> int:
    parser=argparse.ArgumentParser(description=__doc__,formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--input-csv',type=Path,default=Path(CSV_NAME))
    parser.add_argument('--model-file',type=Path,default=Path(BASE_NAME))
    parser.add_argument('--output-dir',type=Path,default=Path('.'))
    parser.add_argument('--no-docx',action='store_true')
    args=parser.parse_args()
    try:
        src,model_path=args.input_csv.resolve(),args.model_file.resolve()
        if not src.is_file() or not model_path.is_file():
            raise FileNotFoundError('The original all-scenario CSV and simulation script are required. Use --input-csv and --model-file or run from the repository root.')
        src_hash,model_hash=file_hash(src),file_hash(model_path)
        df=pd.read_csv(src,float_precision='round_trip'); validate_design(df)
        model=load_model(model_path)
        scenarios,components,checks=analyze(model,df)
        selections,summary,comparison,controls=summarize(model,scenarios,src.parent)
        components=components.merge(scenarios[['scenario_id','feasible_primary','feasible_broader_MV']],
                                    on='scenario_id',validate='many_to_one')
        args.output_dir.mkdir(parents=True,exist_ok=True)
        out=args.output_dir.resolve()
        for name,data in [('compartment_energy_power_scenarios.csv',scenarios),
                          ('compartment_energy_power_components.csv',components),
                          ('compartment_energy_power_selected_strategies.csv',selections),
                          ('compartment_energy_power_summary.csv',summary),
                          ('compartment_energy_power_comparison.csv',comparison)]:
            path=out/name
            if path in (src,model_path): raise ValueError('Input and output paths must differ.')
            data.to_csv(path,index=False,float_format='%.17g')
        checks.update(input_csv_sha256=src_hash,model_sha256=model_hash,dt_seconds=DT,
                      input_csv=src.name,model=model_path.name,feasible_controls=controls,
                      python=platform.python_version(),numpy=np.__version__,pandas=pd.__version__,
                      notes='Reference inspiration-only model. Numerical Euler term is not a physical dissipation. No physiological validation or analysis of combined extensions.')
        if not args.no_docx:
            export_docx(out,selections,summary,comparison,checks)
        if file_hash(src)!=src_hash or file_hash(model_path)!=model_hash:
            raise RuntimeError('An input file unexpectedly changed.')
        (out/'compartment_energy_power_verification.json').write_text(json.dumps(checks,indent=2),encoding='utf-8')
        print('Original model outputs reproduced for all 2,700 scenarios; primary n=324; broader MV n=1,944.')
        print(comparison[['analysis_set','phenotype','MPmax_positive_J_min_change_percent']].to_string(index=False))
        print(f'Wrote outputs to {out}')
        return 0
    except (ValueError,FileNotFoundError,ImportError,RuntimeError,OSError) as e:
        print(f'ERROR: {e}',file=sys.stderr)
        return 1

# The optional document generator is defined below; main is called at file end.

def export_docx(out: Path, selections: pd.DataFrame, summary: pd.DataFrame,
                comparison: pd.DataFrame, checks: dict[str, Any]) -> Path:
    """Create the optional, publication-facing three-page supplementary table."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc=Document()
    sec=doc.sections[0]
    sec.page_width=Inches(11.7); sec.page_height=Inches(8.3)
    sec.top_margin=sec.bottom_margin=Inches(0.45)
    sec.left_margin=sec.right_margin=Inches(0.5)
    sec.header_distance=sec.footer_distance=Inches(0.2)
    normal=doc.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(9)
    normal.paragraph_format.space_after=Pt(4)
    normal.paragraph_format.line_spacing=1.0
    for h in ('Heading 1','Heading 2'):
        doc.styles[h].font.name='Calibri'; doc.styles[h].font.size=Pt(11)
        doc.styles[h].font.color.rgb=RGBColor.from_string('000000')
        doc.styles[h].paragraph_format.space_before=Pt(7)
        doc.styles[h].paragraph_format.space_after=Pt(4)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=footer.add_run('Supplementary Table S7 | '); r.font.size=Pt(8)
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    doc.core_properties.title='Supplementary Table S7. Compartmental inspiratory work and power'
    doc.core_properties.subject='Reference six-compartment model: positive-inflow power and energy-accounting diagnostics'
    doc.core_properties.author=''

    def text(s: str, small=False):
        p=doc.add_paragraph(s)
        if small:
            for r in p.runs:r.font.size=Pt(8.3)
        return p
    def heading(s: str):doc.add_paragraph(s,'Heading 2')
    def title(cont=False):
        s=('Supplementary Table S7 (continued).' if cont else
           'Supplementary Table S7. Compartmental inspiratory work, power, and energy accounting.')
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(5)
        r=p.add_run(s);r.bold=True;r.font.size=Pt(13)
    def table(headers, data, widths, fontsize=8.5):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for col,w in zip(t.columns,widths):col.width=Inches(w)
        for c,w,h in zip(t.rows[0].cells,widths,headers):c.width=Inches(w);c.text=h
        for values in data:
            row=t.add_row()
            for c,w,v in zip(row.cells,widths,values):c.width=Inches(w);c.text=str(v)
        for ir,row in enumerate(t.rows):
            trpr=row._tr.get_or_add_trPr(); tag=OxmlElement('w:cantSplit');trpr.append(tag)
            if ir==0:
                repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
            for ic,c in enumerate(row.cells):
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp=c._tc.get_or_add_tcPr()
                margins=OxmlElement('w:tcMar')
                for side,val in [('top','50'),('bottom','50'),('left','55'),('right','55')]:
                    item=OxmlElement('w:'+side);item.set(qn('w:w'),val);item.set(qn('w:type'),'dxa');margins.append(item)
                cp.append(margins)
                if ir==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EDF0');cp.append(shade)
                elif ir%2==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F7F8F9');cp.append(shade)
                for p in c.paragraphs:
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ic<2 else WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
                    p.paragraph_format.line_spacing=1.0
                    for r in p.runs:r.font.size=Pt(fontsize);r.bold=(ir==0)
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
        return t
    def sel(analysis, strategy, p):
        return selections.loc[(selections.analysis_set==analysis)&(selections.strategy==strategy)&(selections.phenotype==p)].iloc[0]
    def fmt(v,n=3):return f'{float(v):.{n}f}'
    def wave(w):return {'square':'Square','decelerating':'Decelerating','sinusoidal':'Sinusoidal'}[w]
    def strategy_name(s):return 'Minimum MP' if s=='minimum_MP' else 'Balanced'
    def comparison_rows(analysis):
        data=[]
        for p in ORDER:
            for strategy in ('minimum_MP','balanced'):
                r=sel(analysis,strategy,p)
                data.append([LABELS[p],strategy_name(strategy),fmt(r.Ti_s,1),wave(r.waveform),fmt(r.pause_fraction,1),
                             fmt(r.MP_J_min),fmt(r.EII),fmt(r.DCES),fmt(r.Emax_positive_J_per_breath,5),fmt(r.MPmax_positive_J_min,4)])
        return data
    headers=['Configuration','Selection','Ti\n(s)','Waveform','Pause\nfraction','Global MP\n(J/min)','EII','DCES','Emax⁺\n(J/breath)','MPmax⁺\n(J/min)']
    widths=[1.48,1.08,0.43,1.0,0.58,0.75,0.61,0.61,0.88,0.9]
    title()
    text('Reference linear, inspiration-only simulations with volume reset. Original equations, 0.002-s time step, energy definitions, and strategy selections are retained. Eᵢ is the original positive-inflow work per breath; MPᵢ⁺ = RR × Eᵢ. Emax⁺ and MPmax⁺ are maxima across the six compartments, not instantaneous temporal peaks.',True)
    heading('A. Original minimum-MP and balanced selections: primary feasible set (324 scenarios)')
    table(headers,comparison_rows('primary'),widths)
    text('All rows use VT 4 mL/kg PBW, RR 25 breaths/min, MV 7.0 L/min, PBW 70 kg, and PEEP 5 cmH₂O. Relative to minimum MP, balanced selection increases MPmax⁺ by 32.12%, 2.54%, 4.45%, and 5.36% in the listed configuration order, despite reducing EII and DCES.',True)
    heading('B. All six positive-inflow compartmental powers for the primary selections (J/min)')
    data=[]
    for p in ORDER:
        for strategy in ('minimum_MP','balanced'):
            r=sel('primary',strategy,p)
            data.append([LABELS[p],strategy_name(strategy)]+[fmt(r[f'MPplus_{i}_J_min'],4) for i in range(1,7)])
    table(['Configuration','Selection']+[f'MP{i}⁺' for i in range(1,7)],data,[1.48,1.08]+[.97]*6)
    text('Since MPᵢ⁺/ΣⱼMPⱼ⁺ = Eᵢ/ΣⱼEⱼ, REF, EII, and DCES are algebraically invariant to the energy-to-power conversion within each scenario. This identity is not an independent validation. ΣᵢMPᵢ⁺ need not equal global MP because outward compartmental work is excluded from the positive-inflow metric.',True)

    doc.add_page_break();title(True)
    heading('C. Original selections in the broader minute-ventilation set (1,944 scenarios)')
    table(headers,comparison_rows('broader_MV'),widths)
    text('All listed broader-set selections use VT 4 mL/kg, RR 15 breaths/min, and MV 4.2 L/min. Balanced selection increases MPmax⁺ by 72.26%, 8.56%, 4.45%, and 10.58%, respectively. Thus, lower fractional concentration does not establish lower maximum absolute compartmental power in either analysis set.',True)
    heading('D. Separate minimization of maximum energy per breath and maximum power')
    data=[]
    for analysis in ('primary','broader_MV'):
        for p in ORDER:
            er=selections.loc[(selections.analysis_set==analysis)&(selections.phenotype==p)&(selections.strategy=='minimum_maximum_energy')]
            pr=sel(analysis,'minimum_maximum_power',p)
            rrvals=', '.join(str(int(v)) for v in sorted(er.RR_min.unique()))
            powers=f'{er.MPmax_positive_J_min.min():.4f}'
            if len(er)>1:powers+=f'–{er.MPmax_positive_J_min.max():.4f}'
            data.append([('Primary' if analysis=='primary' else 'Broader MV'),LABELS[p],wave(pr.waveform),fmt(er.Emax_positive_J_per_breath.min(),5),
                         str(len(er)),rrvals,powers,str(int(pr.RR_min)),fmt(pr.MPmax_positive_J_min,4),fmt(pr.MP_J_min)])
    table(['Analysis set','Configuration','Selected\nwaveform','Minimum\nEmax⁺\n(J/breath)','Energy\nties (n)','RR among energy\nminima (/min)','MPmax⁺ among\nenergy minima\n(J/min)','RR at power\nminimum (/min)','Minimum\nMPmax⁺\n(J/min)','Global MP at\npower minimum\n(J/min)'],data,[.8,1.38,.97,.82,.51,1.1,1.18,.83,.83,.92],8.1)
    text('All energy and power minima above use VT 4 mL/kg, Ti 1.5 s, and no pause. All tied minima were retained (relative tolerance 10⁻¹⁰; absolute tolerance 10⁻¹²). The power minimum is unique within each configuration and set. In the broader set, the four equal energy minima span RR 15–30/min, but only RR 15/min minimizes power.',True)
    text('Maximum-power minimization coincides with global minimum-MP selection in three configurations. In the resistance-dominant configuration, it selects decelerating rather than square flow. In the primary set this lowers MPmax⁺ from 0.6555 to 0.6409 J/min (−2.22%) while global MP rises from 3.6253 to 3.7209 J/min (+2.64%). These are separate supplementary selection criteria; the original balanced objective is not replaced.',True)
    text('For the broader-set energy minima, the range in MPmax⁺ reflects differences in RR at identical within-breath settings, not uncertainty in the model. Full scenario-level and six-compartment outputs are included in the numerical CSV files.',True)

    doc.add_page_break();title(True)
    heading('E. Energy-accounting diagnostics for the primary selections (all energies in J per breath)')
    data=[]
    for p in ORDER:
        for strategy in ('minimum_MP','balanced'):
            r=sel('primary',strategy,p)
            data.append([LABELS[p],strategy_name(strategy)]+[fmt(r[f],6) for f in (
                'global_energy_J','Eplus_sum_J_per_breath','Eminus_sum_J_per_breath',
                'D_sum_J_per_breath','delta_U_sum_J_per_breath','euler_increment_sum_J_per_breath')])
    table(['Configuration','Selection','Global work\nΣWᵢ','Inward work\nΣEᵢ','Outward work\nΣEᵢ⁻','Branch losses\nΣDᵢ','Elastic storage\nΣΔUᵢ','Euler term\nΣεᵢ'],data,[1.48,1.05]+[1.07]*6,8.4)
    heading('Definitions and numerical implementation')
    text('Let p = Paw − PEEP. Signed branch inlet work is Wᵢ = k∫p qᵢ dt; branch-resistive dissipation during inspiration is Dᵢ = k∫Rᵢ qᵢ² dt; the quadratic elastic-energy change relative to prescribed-PEEP equilibrium is ΔUᵢ = k[Vᵢ(Ti)² − Vᵢ(0)²]/(2Cᵢ). All integrals extend over inspiration including any pause, and k = 0.0980665 J/(cmH₂O·L). Here Vᵢ(0) = 0.')
    text('The continuous RC equation gives Wᵢ = Dᵢ + ΔUᵢ. Outward work magnitude is Eᵢ⁻ = k∫p max(−qᵢ,0) dt, so Eᵢ − Eᵢ⁻ = Wᵢ. The original reported Eᵢ is cumulative inward work, not the signed balance Dᵢ + ΔUᵢ and not parenchymal energy dissipation. With common inlet pressure, ΣWᵢ = Eglobal, whereas ΣEᵢ can exceed Eglobal.')
    text('The diagnostic replay preserves the original forward-Euler updates and pre-update rectangular energy summation. Its exact discrete identity is Wᵢ = Dᵢ + ΔUᵢ − εᵢ, where εᵢ = kΣₙ(qᵢ,ₙ Δt)²/(2Cᵢ) ≥ 0. This term is a discretization residual, not an additional physical loss; it is reported without changing the original MP or Eᵢ. The continuous identity is not imposed exactly on a finite-step Euler trajectory.')
    heading('Scope and verification')
    epsmax=summary.max_euler_increment_relative_to_global_work_percent.max()
    err=checks['max_absolute_identity_residuals']
    text(f'The diagnostics reproduced all 2,700 stored reference scenarios (540 distinct inspirations); the reconstructed primary and broader feasible sets and original balanced scores were checked against their stored CSV files. The maximum branchwise discrete-balance residual was {err["discrete_branch_balance_J"]:.2e} J, and the maximum error in ΣWᵢ = Eglobal was {err["signed_work_sum_J"]:.2e} J. Across each feasible set, the summed Euler term was at most {epsmax:.4f}% of global inspiratory work. The energy-to-power conversion changed REF, EII, and DCES only at floating-point round-off level.')
    text('Dᵢ represents dissipation in the lumped branch resistance, not validated tissue dissipation. ΔUᵢ describes storage relative to the model pressure-volume reference, not absolute lung elastic energy. MPᵢ⁺ describes inspiratory positive-inflow work per minute, not a whole-cycle dissipation rate. The analysis is not combined with finite-expiration, grid-refinement, or PEEP-offset extensions. No new patient measurements or injury validation are implied.',True)
    text('Abbreviations: MP, global inspiratory mechanical power above prescribed PEEP; REF, regional energy fraction; EII, Energy Inequality Index; DCES, Dominant Compartment Energy Share; VT, tidal volume; RR, respiratory rate; Ti, total inspiratory time; MV, minute ventilation; PBW, predicted body weight; PEEP, positive end-expiratory pressure. EII and DCES are dimensionless. Values were computed from unrounded outputs and rounded only for presentation.',True)
    path=out/'Supplementary_Table_S7_compartment_energy_power.docx';doc.save(path);return path


if __name__=='__main__':
    raise SystemExit(main())
