#!/usr/bin/env python3
"""Objective dependence and secondary scalarization sensitivity (Table S8).

Post-processing only: reads the original 2,700-scenario, inspiration-only CSV.
No model trajectories, original results, objectives, or Pareto flags are changed.

Run from the repository root:
    python assess_objective_sensitivity.py --output-dir . --no-docx
Omit --no-docx to generate Supplementary Table S8 using python-docx.

The named controls use the original min-max bounds within each configuration
and feasible set. A strictly positive weight grid samples all (i,j,k)/20 with
positive integer i+j+k=20 (171 triples), plus equal weights (1/3,1/3,1/3).
This is a post hoc sensitivity diagnostic, NOT weight validation, estimation
of preference probabilities, or a new recommended selection criterion.
"""
from __future__ import annotations
import argparse
import csv
import hashlib
import itertools
import json
from pathlib import Path
import platform
import sys
from typing import Any
import numpy as np

RTOL, ATOL = 1e-10, 1e-12
ORDER = ('compliance_dominant','resistance_dominant','mixed','severe_mixed')
LABELS = dict(zip(ORDER, ('Compliance-dominant','Resistance-dominant','Mixed','Severe mixed')))
KEYS = ('phenotype','VT_ml_per_kg','RR_min','Ti_s','waveform','pause_fraction','PEEP_cmH2O')
METRICS = ('MP_J_min','EII','DCES')
POWER = 'MPmax_positive_J_min'
CONTROLS = (
    ('equal_three', (1/3,1/3,1/3)),
    ('omit_DCES', (0.5,0.5,0.0)),
    ('omit_EII', (0.5,0.0,0.5)),
    ('half_global_half_distribution', (0.5,0.25,0.25)),
)
CONTROL_LABELS = {'equal_three':'Equal three', 'omit_DCES':'Omit DCES',
                  'omit_EII':'Omit EII','half_global_half_distribution':'Half global / half regional'}
BASE_CSV = 'MVtarget_primary_all_simulation_results.csv'


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_reference(path: Path) -> list[dict[str, Any]]:
    with path.open(newline='',encoding='utf-8-sig') as f:
        raw=list(csv.DictReader(f))
    if len(raw)!=2700:
        raise ValueError('Expected the original all-scenario CSV with 2,700 rows.')
    needed=set(KEYS)|set(METRICS)|{'PBW_kg','VT_ml','minute_ventilation_L_min',
        'Ppeak_cmH2O','global_energy_J'}|{f'comp_energy_J_{i}' for i in range(1,7)}
    if missing:=needed-set(raw[0]):
        raise ValueError(f'Missing reference columns: {sorted(missing)}')
    records=[]
    for i, row in enumerate(raw,1):
        r={k:(row[k] if k in ('phenotype','waveform') else float(row[k])) for k in sorted(needed)}
        if not all(np.isfinite(v) for v in r.values() if isinstance(v,float)):
            raise ValueError(f'Nonfinite numeric value in row {i}')
        r['scenario_id']=i
        e=np.array([r[f'comp_energy_J_{j}'] for j in range(1,7)],float)
        if np.any(e<0) or e.sum()<=0 or r['RR_min']<=0:
            raise ValueError('Nonpositive energy sum or respiratory rate.')
        ref=e/e.sum()
        for observed,expected,label in (
                (r['EII'],np.std(ref)/np.mean(ref),'EII'),
                (r['DCES'],max(ref),'DCES'),
                (r['MP_J_min'],r['global_energy_J']*r['RR_min'],'global MP'),
                (r['VT_ml'],r['PBW_kg']*r['VT_ml_per_kg'],'VT'),
                (r['minute_ventilation_L_min'],r['VT_ml']*r['RR_min']/1000,'MV')):
            if not np.isclose(observed,expected,rtol=RTOL,atol=ATOL):
                raise ValueError(f'{label} inconsistency in reference row {i}')
        r['Eplus_sum_J_per_breath']=float(e.sum())
        r['MPplus_sum_J_min']=float(r['RR_min']*e.sum())
        r[POWER]=float(np.max(r['RR_min']*e))
        records.append(r)
    expected=set(itertools.product(ORDER,(4,5,6,7,8),(10,15,20,25,30),
                                  (0.6,1.0,1.5),('square','decelerating','sinusoidal'),(0,0.1,0.2),(5.,)))
    actual=[tuple(r[k] for k in KEYS) for r in records]
    if set(actual)!=expected or len(set(actual))!=2700:
        raise ValueError('Input does not match the original 2,700-scenario grid.')
    if any(not np.isclose(r['PBW_kg'],70,rtol=0,atol=ATOL) for r in records):
        raise ValueError('Expected PBW 70 kg.')
    return records


def tied_indices(scores: np.ndarray) -> np.ndarray:
    return np.flatnonzero(np.isclose(scores,float(np.min(scores)),rtol=RTOL,atol=ATOL))


def dump_csv(path: Path, records: list[dict[str,Any]]) -> None:
    if not records:
        raise ValueError(f'No records for {path.name}')
    keys=list(dict.fromkeys(k for r in records for k in r))
    with path.open('w',newline='',encoding='utf-8') as f:
        writer=csv.DictWriter(f,fieldnames=keys);writer.writeheader()
        for r in records:
            writer.writerow({k:(format(v,'.17g') if isinstance(v,(float,np.floating)) else v)
                             for k,v in r.items()})


def analyze(records: list[dict[str,Any]], source_dir: Path):
    selections=[];weight_rows=[];summaries=[];dominance=[];verification={}
    weights=[(i/20,j/20,(20-i-j)/20) for i in range(1,19) for j in range(1,20-i)]
    weights.append((1/3,1/3,1/3))
    if len(weights)!=172 or not np.all(np.array(weights)>0):
        raise ValueError('Unexpected positive weight-grid definition.')
    max_identity=0.0
    for analysis_name,lo,hi,n_expected,strat_name in (
            ('primary',6.5,7.5,324,'MVtarget_primary_strategy_summary.csv'),
            ('broader_MV',4.,12.,1944,'MVtarget_sensitivity_strategy_summary.csv')):
        subset=[r for r in records if r['Ppeak_cmH2O']<=35 and lo<=r['minute_ventilation_L_min']<=hi]
        if len(subset)!=n_expected:raise ValueError(f'Unexpected {analysis_name} feasible count.')
        stored_path=source_dir/strat_name
        stored=[]
        if stored_path.is_file():
            with stored_path.open(newline='',encoding='utf-8-sig') as f:stored=list(csv.DictReader(f))
        for phenotype in ORDER:
            group=[r for r in subset if r['phenotype']==phenotype]
            x=np.array([[r[k] for k in METRICS] for r in group],float)
            bounds_min,bounds_max=x.min(0),x.max(0)
            span=bounds_max-bounds_min
            if np.any(span<=0):raise ValueError('A normalization range is zero.')
            z=(x-bounds_min)/span
            original_score=z.sum(axis=1)
            baseline_i=int(np.argmin(original_score));baseline=group[baseline_i]
            if stored:
                sb=[r for r in stored if r['phenotype']==phenotype and r['strategy']=='balanced']
                if len(sb)!=1:raise ValueError('Stored balanced selection is not unique.')
                for k in KEYS:
                    if k in ('phenotype','waveform'):
                        ok=sb[0][k]==baseline[k]
                    else:ok=np.isclose(float(sb[0][k]),baseline[k],rtol=RTOL,atol=ATOL)
                    if not ok:raise ValueError(f'Original balanced setting mismatch: {analysis_name}/{phenotype}/{k}')
            for criterion,w in CONTROLS:
                scores=z@np.array(w)
                ids=tied_indices(scores)
                for j in ids:
                    selections.append(dict(analysis_set=analysis_name,criterion=criterion,
                        weight_MP=w[0],weight_EII=w[1],weight_DCES=w[2],
                        tied_minima_count=len(ids),score=float(scores[j]),
                        original_balanced_score=float(original_score[j]),**group[j]))
            for criterion,metric in (('minimum_MP','MP_J_min'),('minimum_EII','EII'),
                                     ('minimum_DCES','DCES'),('minimum_maximum_power',POWER)):
                scores=np.array([r[metric] for r in group])
                ids=tied_indices(scores)
                for j in ids:
                    selections.append(dict(analysis_set=analysis_name,criterion=criterion,
                        tied_minima_count=len(ids),score=float(scores[j]),
                        original_balanced_score=float(original_score[j]),**group[j]))
            unique=set();tied_weights=0
            for wid,w in enumerate(weights,1):
                scores=z@np.array(w);ids=tied_indices(scores)
                tied_weights+=int(len(ids)>1)
                for j in ids:
                    unique.add(int(j))
                    weight_rows.append(dict(analysis_set=analysis_name,weight_id=wid,
                        weight_MP=w[0],weight_EII=w[1],weight_DCES=w[2],
                        tied_minima_count=len(ids),score=float(scores[j]),**group[j]))
            # A direct comparison to the original balanced setting under the
            # explicitly different objective vector (MP, maximum power, EII).
            # This is not a recomputation or replacement of the original Pareto front.
            alt=np.array([[r[k] for k in ('MP_J_min',POWER,'EII')] for r in group])
            close=np.isclose(alt,alt[baseline_i],rtol=RTOL,atol=ATOL)
            dominators=np.flatnonzero(np.all((alt<alt[baseline_i])|close,axis=1)&
                                     np.any((alt<alt[baseline_i])&~close,axis=1))
            for j in dominators:
                r=group[j]
                dominance.append(dict(analysis_set=analysis_name,phenotype=phenotype,
                    reference_scenario_id=baseline['scenario_id'],candidate_scenario_id=r['scenario_id'],
                    reference_waveform=baseline['waveform'],candidate_waveform=r['waveform'],
                    reference_pause_fraction=baseline['pause_fraction'],candidate_pause_fraction=r['pause_fraction'],
                    **{f'{prefix}_{k}':row[k] for prefix,row in [('reference',baseline),('candidate',r)]
                       for k in ('VT_ml_per_kg','RR_min','Ti_s','MP_J_min',POWER,'EII','DCES')}))
            summary=dict(analysis_set=analysis_name,phenotype=phenotype,n_feasible=len(group),
                sampled_positive_weight_triples=len(weights),distinct_selected_settings=len(unique),
                weight_triples_with_tied_minima=tied_weights,original_balanced_scenario_id=baseline['scenario_id'],
                original_balanced_MP_J_min=baseline['MP_J_min'],original_balanced_MPmax_J_min=baseline[POWER],
                original_balanced_EII=baseline['EII'],original_balanced_DCES=baseline['DCES'],
                balanced_dominators_for_MP_MPmax_EII=len(dominators))
            for k in ('MP_J_min',POWER,'EII','DCES'):
                values=[group[j][k] for j in unique]
                summary[f'selected_{k}_min']=min(values);summary[f'selected_{k}_max']=max(values)
            for col,k in enumerate(METRICS):
                summary[f'{k}_normalization_min']=float(bounds_min[col])
                summary[f'{k}_normalization_max']=float(bounds_max[col])
            for r in group:
                max_identity=max(max_identity,abs(r[POWER]-r['DCES']*r['MPplus_sum_J_min']))
            summaries.append(summary)
        verification[f'{analysis_name}_n']=len(subset)
        verification[f'{analysis_name}_stored_balanced_check']=('Matched all four original settings.'
            if stored else 'Recomputed from reference CSV; separate strategy-summary CSV not supplied.')
    if max_identity>1e-10:raise ValueError('Maximum-power/share identity check failed.')
    verification.update(max_absolute_MPmax_share_identity_error_J_min=max_identity,
        sampled_positive_weight_triples=172,source_scenarios=2700,
        new_dynamics_simulated=False,original_files_modified=False,
        original_Pareto_fronts_recomputed=False,
        tolerance_note='Minimizer ties and direct objective dominance: rtol=1e-10, atol=1e-12.',
        scope='Reference inspiration-only sets; no combination with finite expiration, PEEP offsets, or refined setting grids.')
    return selections,weight_rows,summaries,dominance,verification


def export_docx(out: Path, selections, summaries, dominance) -> Path:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document();s=doc.sections[0]
    s.page_width=Inches(11.7);s.page_height=Inches(8.3)
    s.left_margin=s.right_margin=Inches(.5);s.top_margin=s.bottom_margin=Inches(.45)
    s.footer_distance=Inches(.2)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.0
    for sty in ('Heading 1','Heading 2'):
        doc.styles[sty].font.name='Calibri';doc.styles[sty].font.size=Pt(11)
        doc.styles[sty].font.color.rgb=RGBColor.from_string('000000')
        doc.styles[sty].paragraph_format.space_before=Pt(8)
        doc.styles[sty].paragraph_format.space_after=Pt(5)
    f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    f.add_run('Supplementary Table S8 | ').font.size=Pt(8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
    doc.core_properties.title='Supplementary Table S8. Objective dependence and scalarization sensitivity'
    doc.core_properties.author=''
    def title(continued=False):
        p=doc.add_paragraph();p.paragraph_format.page_break_before=continued
        r=p.add_run('Supplementary Table S8 (continued).' if continued else
            'Supplementary Table S8. Objective dependence and scalarization sensitivity.');r.bold=True;r.font.size=Pt(13)
    def text(content, small=False):
        p=doc.add_paragraph(content)
        if small:
            for r in p.runs:r.font.size=Pt(9)
        return p
    def heading(content):doc.add_paragraph(content,'Heading 2')
    def table(headers, values, widths, size=9):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for col,width in zip(t.columns,widths):col.width=Inches(width)
        for c,width,h in zip(t.rows[0].cells,widths,headers):c.width=Inches(width);c.text=h
        for vals in values:
            row=t.add_row()
            for c,width,v in zip(row.cells,widths,vals):c.width=Inches(width);c.text=str(v)
        for i,row in enumerate(t.rows):
            trp=row._tr.get_or_add_trPr();trp.append(OxmlElement('w:cantSplit'))
            if i==0:trp.append(OxmlElement('w:tblHeader'))
            for j,c in enumerate(row.cells):
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=c._tc.get_or_add_tcPr()
                if i==0 or i%2==0:
                    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E9EDF0' if i==0 else 'F6F7F8');pr.append(fill)
                mar=OxmlElement('w:tcMar')
                for side,value in [('top','55'),('bottom','55'),('left','55'),('right','55')]:
                    v=OxmlElement('w:'+side);v.set(qn('w:w'),value);v.set(qn('w:type'),'dxa');mar.append(v)
                pr.append(mar)
                for p in c.paragraphs:
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j<3 else WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
                    for r in p.runs:r.font.size=Pt(size);r.bold=(i==0)
        return t
    def control_table(analysis):
        data=[]
        for ph in ORDER:
            for cname,_ in CONTROLS:
                choices=[r for r in selections if r['analysis_set']==analysis and r['phenotype']==ph and r['criterion']==cname]
                if len(choices)!=1:raise ValueError('Document control table expects unique control selections; inspect CSV ties.')
                r=choices[0]
                data.append([LABELS[ph],CONTROL_LABELS[cname],r['waveform'].capitalize(),
                    f"{r['Ti_s']:.1f}",f"{r['pause_fraction']:.1f}",f"{r['MP_J_min']:.4f}",
                    f"{r[POWER]:.4f}",f"{r['EII']:.6f}",f"{r['DCES']:.6f}"])
        table(['Configuration','Score / weights','Waveform','Ti\n(s)','Pause','Global MP\n(J/min)',
               'MPmax⁺\n(J/min)','EII','DCES'],data,[1.42,1.55,1.1,.46,.52,.92,.92,.95,.95],9)
    title()
    text('Secondary, post hoc analysis of the stored reference inspiration-only results. No new respiratory-model simulations were performed. The primary and broader feasible sets contain 324 and 1,944 scenarios, respectively. The original model, numerical outputs, balanced selections, and Pareto flags are not replaced.',True)
    heading('A. Named scalarization controls: primary feasible set')
    control_table('primary')
    text('All rows: VT 4 mL/kg PBW, RR 25 breaths/min, MV 7.0 L/min; PBW 70 kg and PEEP 5 cmH₂O. All listed control minima are unique within the specified tolerance.',True)
    heading('Score definitions')
    text('Weight order is (MP, EII, DCES), applied to the original within-configuration, within-set min–max-normalized objectives. Equal three: (1/3, 1/3, 1/3); omit DCES: (1/2, 1/2, 0); omit EII: (1/2, 0, 1/2); half global / half regional: (1/2, 1/4, 1/4). Equal three is the original balanced score divided by three and has the same minimizer. The controls are sensitivity checks, not recommendations for selecting physiological weights.',True)
    text('In the severe mixed primary configuration, omitting DCES selects decelerating flow with pause 0.1 rather than sinusoidal flow without a pause. Global MP, MPmax⁺, and EII are all lower, but DCES is slightly higher. Thus, the original balanced selection is dominated under the different objective vector (MP, MPmax⁺, EII), not under the original vector (MP, EII, DCES). This distinction does not change the previously reported original Pareto analysis.',True)
    title(True)
    heading('B. Named scalarization controls: broader minute-ventilation feasible set')
    control_table('broader_MV')
    text('All rows: VT 4 mL/kg PBW, RR 15 breaths/min, MV 4.2 L/min; PBW 70 kg and PEEP 5 cmH₂O. The same four weight definitions are used, with normalization bounds calculated separately for this broader feasible set. Within a set, bounds are held fixed across all weight vectors.',True)
    text('Omitting EII changes the selected settings in the compliance-dominant, resistance-dominant, and severe mixed configurations. Assigning half the total weight to global MP changes the resistance-dominant and severe mixed selections. Omitting DCES does not change the broader-set equal-three selections. These findings concern selection dependence, not biological validity or clinical performance.',True)
    heading('C. Direct comparison under an absolute-power objective')
    if len(dominance)!=1:raise ValueError('Unexpected number of direct dominance comparisons for document summary.')
    d=dominance[0]
    data=[]
    for label,prefix in [('Original balanced','reference'),('Alternative feasible setting','candidate')]:
        data.append([label,d[f'{prefix}_waveform'].capitalize(),f"{d[f'{prefix}_pause_fraction']:.1f}"]+
                    [f"{d[f'{prefix}_{k}']:.6f}" for k in ('MP_J_min',POWER,'EII','DCES')])
    table(['Severe mixed, primary','Waveform','Pause','Global MP','MPmax⁺','EII','DCES'],data,
          [1.9,1.2,.6,1.1,1.1,1.1,1.1],9)
    text('Power units are J/min. Both settings use VT 4 mL/kg, RR 25/min, and Ti 1.5 s. The alternative has lower values for all three compared objectives (MP, MPmax⁺, EII), while its larger DCES explains why it does not dominate the original balanced selection when DCES is retained. All feasible candidates were checked; this was the only original balanced selection with a dominator under this alternative three-objective comparison.',True)
    title(True)
    heading('D. Sampled strictly positive weights')
    data=[]
    for r in summaries:
        data.append(['Primary' if r['analysis_set']=='primary' else 'Broader MV',LABELS[r['phenotype']],
            str(r['distinct_selected_settings']),
            f"{r['selected_MP_J_min_min']:.3f}–{r['selected_MP_J_min_max']:.3f}",
            f"{r['selected_'+POWER+'_min']:.3f}–{r['selected_'+POWER+'_max']:.3f}",
            f"{r['selected_EII_min']:.4f}–{r['selected_EII_max']:.4f}",
            f"{r['selected_DCES_min']:.4f}–{r['selected_DCES_max']:.4f}"])
    table(['Analysis set','Configuration','Distinct\nsettings','Global MP\nrange (J/min)',
           'MPmax⁺\nrange (J/min)','EII range','DCES range'],data,[.95,1.5,.7,1.4,1.4,1.35,1.35],9)
    text('Each configuration and feasible set was evaluated using 172 weight triples: all strictly positive multiples of 0.05 summing to one (171 triples), plus (1/3, 1/3, 1/3). All tied minima were retained; no ties occurred for these sampled positive weights. The counts and ranges describe this chosen finite weight sample, not all possible weights, sampling uncertainty, or probabilities of clinical preferences.',True)
    heading('Dependence of the objectives')
    text('Let fᵢ = Eᵢ/ΣⱼEⱼ. EII summarizes dispersion across all six fractions, whereas DCES = maxᵢ(fᵢ) reports the largest fraction. Both summarize the same normalized distribution and are not independent physiological measurements. They need not rank candidate settings identically. With the population-SD convention, EII² = 6Σᵢfᵢ² − 1. Maximum positive-inflow power satisfies MPmax⁺ = DCES × RRΣᵢEᵢ, where RRΣᵢEᵢ need not equal global MP because it includes inward transfers during redistribution.')
    heading('Methods, checks, and scope')
    text('All computations use unrounded stored energies and metrics. Original feasibility filters are reapplied (Ppeak ≤35 cmH₂O; MV 6.5–7.5 or 4.0–12.0 L/min). Min–max bounds are recalculated separately within each configuration and feasible set, then held fixed across weight vectors. Tied minima and direct objective comparisons use relative tolerance 10⁻¹⁰ and absolute tolerance 10⁻¹². Normalization bounds and all selected settings are supplied in the numerical outputs. Original balanced settings are checked against stored strategy summaries when those files are present.')
    text('This analysis does not establish an optimal weighting, independent injury prediction, physiological superiority, or structural robustness. It does not recompute the original Pareto fronts and is not combined with the passive-expiration, PEEP-offset, or refined-grid extensions. The original equal-weight score is a secondary illustration; it adds a preference rule, not a fourth independently measured quantity.',True)
    text('Abbreviations: MP, global inspiratory mechanical power above PEEP; MPmax⁺, maximum compartmental positive-inflow power (maximum across compartments, not time); EII, Energy Inequality Index; DCES, Dominant Compartment Energy Share; VT, tidal volume; RR, respiratory rate; Ti, total inspiratory time; MV, minute ventilation; PBW, predicted body weight; PEEP, positive end-expiratory pressure. EII and DCES are dimensionless.',True)
    path=out/'Supplementary_Table_S8_objective_sensitivity.docx';doc.save(path);return path


def main() -> int:
    parser=argparse.ArgumentParser(description=__doc__,formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--input-csv',type=Path,default=Path(BASE_CSV))
    parser.add_argument('--output-dir',type=Path,default=Path('.'))
    parser.add_argument('--no-docx',action='store_true')
    args=parser.parse_args()
    try:
        src=args.input_csv.resolve()
        if not src.is_file():raise FileNotFoundError(f'Original all-scenario CSV not found: {src}')
        original_hash=sha(src);records=read_reference(src)
        sel,wrows,summary,dom,checks=analyze(records,src.parent)
        out=args.output_dir.resolve();out.mkdir(parents=True,exist_ok=True)
        for name,values in [('objective_sensitivity_selected_strategies.csv',sel),
                            ('objective_weight_selections.csv',wrows),
                            ('objective_sensitivity_summary.csv',summary),
                            ('objective_dominance_comparison.csv',dom)]:
            if (out/name).resolve()==src:raise ValueError('Refusing to overwrite the source CSV.')
            dump_csv(out/name,values)
        if not args.no_docx:export_docx(out,sel,summary,dom)
        if sha(src)!=original_hash:raise RuntimeError('Source CSV changed unexpectedly.')
        checks.update(input_csv=src.name,input_sha256=original_hash,python=platform.python_version(),numpy=np.__version__,
                      selected_strategy_rows=len(sel),positive_weight_selected_rows=len(wrows),
                      output_summary_rows=len(summary),direct_dominance_comparisons=len(dom))
        (out/'objective_sensitivity_verification.json').write_text(json.dumps(checks,indent=2),encoding='utf-8')
        print(json.dumps(checks,indent=2))
        for r in summary:print(r['analysis_set'],r['phenotype'],'distinct settings:',r['distinct_selected_settings'])
        return 0
    except (ValueError,FileNotFoundError,ImportError,OSError,RuntimeError) as exc:
        print(f'ERROR: {exc}',file=sys.stderr);return 1

if __name__=='__main__':
    raise SystemExit(main())
