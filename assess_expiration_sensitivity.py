#!/usr/bin/env python3
"""Focused sensitivity analysis of passive expiration and breath-to-breath history.

This ADDITIONAL analysis does not overwrite or silently replace the original
inspiration-only simulations. It re-evaluates the original 324 primary scenarios.

Assumptions introduced only here:
- passive expiration into the prescribed PEEP (5 cmH2O);
- identical linear inspiratory and expiratory branch resistances;
- no shared circuit/expiratory-valve resistance, flow limitation, airway closure,
  nonlinear mechanics, patient effort, recruitment, or gas exchange;
- Te = 60/RR - Ti; Ti includes the end-inspiratory pause;
- exact exponential expiratory emptying; compartment volumes retained between
  breaths until a periodic state is reached.

Inspiration retains the ORIGINAL explicit-Euler step, flow waveforms, and
above-set-PEEP global/positive-inflow regional energy definitions. Expiratory
resistive dissipation is reported separately, NOT added to inspiratory work.
PEEPi_eq = sum(Vee_i)/sum(C_i) is an analytical fully equilibrated pressure
increment, not a clinical measurement or the pressure at the open airway outlet.

Run alongside the original model and CSV:
  python assess_expiration_sensitivity.py --output-dir expiration_sensitivity
Dependencies: numpy, pandas; python-docx only for the optional table export.
Use --no-docx to export numerical results only.
"""
from __future__ import annotations

import argparse
from dataclasses import dataclass
import hashlib
import importlib.util
import json
from pathlib import Path
import platform
from typing import Any

import numpy as np
import pandas as pd

MODEL_NAME = 'multiobjective_ventilation_optimization_scaled_CR_MVtarget.py'
INPUT_NAME = 'MVtarget_primary_feasible_simulation_results.csv'
ORDER = ('compliance_dominant', 'resistance_dominant', 'mixed', 'severe_mixed')
LABEL = {'compliance_dominant': 'Compliance-dominant',
         'resistance_dominant': 'Resistance-dominant',
         'mixed': 'Mixed', 'severe_mixed': 'Severe mixed'}
DT = 0.002
PERIODIC_ATOL_L = 1e-12
SELECTION_RTOL, SELECTION_ATOL = 1e-10, 1e-12
CONVERSION = 0.0980665


@dataclass
class Inspiration:
    end_volume: np.ndarray
    global_energy: float
    positive_energy: np.ndarray
    signed_energy: np.ndarray
    resistive_energy: float
    euler_energy_increment: float
    peak_above_peep: float
    ref: np.ndarray
    eii: float
    dces: float


def inspiration(flow: np.ndarray, C: np.ndarray, R: np.ndarray,
                initial_volume: np.ndarray, dt: float = DT) -> Inspiration:
    """Original left-sum/Euler update order, allowing a nonzero initial state."""
    V = np.asarray(initial_volume, dtype=float).copy()
    inv_R, inv_CR = 1.0 / R, 1.0 / (C * R)
    denominator = np.sum(inv_R)
    global_work = resistive_work = euler_increment = 0.0
    positive_work = np.zeros(len(C))
    signed_work = np.zeros(len(C))
    peak = -np.inf
    for prescribed_flow in flow:
        pressure = (prescribed_flow + np.sum(V * inv_CR)) / denominator
        q = pressure * inv_R - V * inv_CR
        global_work += pressure * prescribed_flow * dt
        positive_work += pressure * np.maximum(q, 0.0) * dt
        signed_work += pressure * q * dt
        resistive_work += np.sum(R * q * q) * dt
        # Exactly the elastic-energy increment introduced by the Euler update.
        euler_increment += np.sum((q * dt) ** 2 / (2.0 * C))
        V += q * dt
        peak = max(peak, pressure)
    ref = positive_work / np.sum(positive_work)
    return Inspiration(V, float(global_work * CONVERSION), positive_work * CONVERSION,
                       signed_work * CONVERSION, float(resistive_work * CONVERSION),
                       float(euler_increment * CONVERSION), float(peak), ref,
                       float(np.std(ref) / np.mean(ref)), float(np.max(ref)))


def transition(C: np.ndarray, R: np.ndarray, n_steps: int, dt: float = DT) -> np.ndarray:
    """Homogeneous map of the SAME discrete inspiratory Euler system."""
    inv_R, inv_CR = 1.0 / R, 1.0 / (C * R)
    B = np.outer(inv_R, inv_CR) / np.sum(inv_R) - np.diag(inv_CR)
    return np.linalg.matrix_power(np.eye(len(C)) + dt * B, n_steps)


def periodic_volume(phi: np.ndarray, zero_reset_end: np.ndarray,
                    decay: np.ndarray) -> tuple[np.ndarray, int]:
    """Iterate the affine full-cycle map; avoid resetting volume after expiration."""
    current = np.zeros(len(zero_reset_end))
    for cycle in range(1, 10001):
        following = decay * (phi @ current + zero_reset_end)
        if np.max(np.abs(following - current)) < PERIODIC_ATOL_L:
            return following, cycle
        current = following
    raise RuntimeError('Periodic state did not converge in 10000 cycle-map iterations.')


def load_model(path: Path) -> Any:
    spec = importlib.util.spec_from_file_location('original_ventilation_model', path)
    if spec is None or spec.loader is None:
        raise ImportError(f'Cannot load model: {path}')
    model = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(model)
    for required in ('get_phenotypes', 'generate_flow_waveform'):
        if not hasattr(model, required):
            raise ValueError(f'Model lacks required function: {required}')
    return model


def validate_input(df: pd.DataFrame) -> None:
    required = ['phenotype', 'VT_ml', 'VT_ml_per_kg', 'RR_min', 'Ti_s', 'waveform',
                'pause_fraction', 'PEEP_cmH2O', 'minute_ventilation_L_min',
                'Ppeak_cmH2O', 'MP_J_min', 'EII', 'DCES', 'balanced_score']
    missing = sorted(set(required) - set(df))
    if missing:
        raise ValueError(f'Missing input columns: {missing}')
    if not df[required].notna().all().all():
        raise ValueError('Missing values in required input columns.')
    if len(df) != 324 or df.groupby('phenotype').size().to_dict() != dict.fromkeys(ORDER, 81):
        raise ValueError('This sensitivity analysis expects the original primary 324-scenario set.')
    if not (np.isclose(df.PEEP_cmH2O, 5.0).all()
            and df.minute_ventilation_L_min.between(6.5, 7.5).all()
            and (df.Ppeak_cmH2O <= 35.0).all()):
        raise ValueError('Input differs from the original primary feasibility conditions.')
    if not ((60.0 / df.RR_min - df.Ti_s) > 0).all():
        raise ValueError('Nonpositive expiratory time in the input.')


def add_periodic_scores(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df['periodic_feasible'] = ((df.periodic_Ppeak_cmH2O <= 35.0)
                               & df.minute_ventilation_L_min.between(6.5, 7.5))
    for metric in ('MP_J_min', 'EII', 'DCES'):
        df['periodic_' + metric + '_norm'] = np.nan
    for phenotype in ORDER:
        ix = df.index[df.phenotype.eq(phenotype) & df.periodic_feasible]
        if len(ix) == 0:
            raise ValueError(f'No periodically feasible scenarios: {phenotype}')
        for metric in ('MP_J_min', 'EII', 'DCES'):
            value = df.loc[ix, 'periodic_' + metric]
            lower, upper = value.min(), value.max()
            df.loc[ix, 'periodic_' + metric + '_norm'] = (
                0.0 if np.isclose(upper, lower) else (value - lower) / (upper - lower))
    df['periodic_balanced_score'] = sum(df['periodic_' + m + '_norm']
                                      for m in ('MP_J_min', 'EII', 'DCES'))
    return df


def select_strategies(df: pd.DataFrame) -> pd.DataFrame:
    selections = []
    for phenotype in ORDER:
        sub = df[df.phenotype.eq(phenotype)]
        for basis in ('original', 'periodic'):
            candidates = sub if basis == 'original' else sub[sub.periodic_feasible]
            for strategy, field in (('minimum_MP', 'MP_J_min'), ('balanced', 'balanced_score')):
                column = field if basis == 'original' else 'periodic_' + field
                minimum = candidates[column].min()
                tied = candidates[np.isclose(candidates[column], minimum,
                                              rtol=SELECTION_RTOL, atol=SELECTION_ATOL)].copy()
                tied.insert(1, 'selection_basis', basis)
                tied.insert(2, 'strategy', strategy)
                tied.insert(3, 'n_tied_minima', len(tied))
                selections.append(tied)
    return pd.concat(selections, ignore_index=True)


def configuration_summary(df: pd.DataFrame, selected: pd.DataFrame) -> pd.DataFrame:
    result = []
    for phenotype in ORDER:
        sub = df[df.phenotype.eq(phenotype)]
        row = {'phenotype': phenotype, 'n_original_primary': len(sub),
               'n_periodic_feasible': int(sub.periodic_feasible.sum())}
        for field in ('Te_s', 'MP_increase_percent', 'excess_end_expiratory_volume_mL',
                      'PEEPi_eq_cmH2O', 'periodic_Ppeak_cmH2O', 'periodic_MP_J_min',
                      'periodic_EII', 'periodic_DCES', 'expiratory_resistive_energy_J'):
            row[field + '_min'] = float(sub[field].min())
            row[field + '_max'] = float(sub[field].max())
        for strategy in ('minimum_MP', 'balanced'):
            s = selected[selected.phenotype.eq(phenotype) & selected.strategy.eq(strategy)]
            original = set(s.loc[s.selection_basis.eq('original'), 'scenario_id'])
            periodic = set(s.loc[s.selection_basis.eq('periodic'), 'scenario_id'])
            row[strategy + '_selection_set_unchanged'] = original == periodic
        result.append(row)
    return pd.DataFrame(result)


def self_tests() -> dict[str, float]:
    """Check homogeneous analytical solution and long-expiration limit."""
    C, R = np.full(6, .04 / 6), np.full(6, 48.)
    ti, te, vt = 1.0, .6, .28
    Q = np.full(round(ti / DT), vt / ti)
    zero = inspiration(Q, C, R, np.zeros(6))
    phi = transition(C, R, len(Q))
    decay = np.exp(-te / (C * R))
    V, _ = periodic_volume(phi, zero.end_volume, decay)
    expected_volume = vt * decay[0] / (1.0 - decay[0])
    out = inspiration(Q, C, R, V)
    expected_extra_energy = CONVERSION * expected_volume / np.sum(C) * vt
    err_v = abs(np.sum(V) - expected_volume)
    err_e = abs((out.global_energy - zero.global_energy) - expected_extra_energy)
    long_V, _ = periodic_volume(phi, zero.end_volume, np.exp(-50.0 * np.ones(6)))
    if max(err_v, err_e, np.max(np.abs(long_V))) > 1e-10:
        raise ValueError('Analytical self-test failed.')
    return {'homogeneous_volume_error_L': float(err_v),
            'homogeneous_work_error_J': float(err_e),
            'long_expiration_residual_volume_L': float(np.sum(long_V))}


def analyse(df: pd.DataFrame, model: Any) -> tuple[pd.DataFrame, dict[str, Any]]:
    mechanics = model.get_phenotypes()
    rows = []
    max_baseline_error = max_direct_error = 0.0
    direct_checked = set()
    for idx, r in df.reset_index(drop=True).iterrows():
        C, R = mechanics[r.phenotype]['C'], mechanics[r.phenotype]['R']
        flow = model.generate_flow_waveform(r.VT_ml / 1000.0, r.Ti_s,
                                             r.pause_fraction, r.waveform, dt=DT)
        zero = inspiration(flow, C, R, np.zeros(len(C)))
        reproduced = np.array([zero.global_energy * r.RR_min, zero.eii, zero.dces,
                               zero.peak_above_peep + r.PEEP_cmH2O])
        expected = np.array([r.MP_J_min, r.EII, r.DCES, r.Ppeak_cmH2O])
        if not np.allclose(reproduced, expected, rtol=1e-11, atol=1e-12):
            raise ValueError(f'Original data not reproduced, input row {idx + 1}.')
        max_baseline_error = max(max_baseline_error, float(np.max(np.abs(reproduced-expected))))
        te = 60.0 / r.RR_min - r.Ti_s
        decay = np.exp(-te / (R * C))
        phi = transition(C, R, len(flow))
        Vee, iterations = periodic_volume(phi, zero.end_volume, decay)
        per = inspiration(flow, C, R, Vee)
        next_Vee = decay * per.end_volume
        closure = float(np.max(np.abs(next_Vee - Vee)))
        if closure > 1e-10:
            raise ValueError(f'Actual full-cycle integration did not close, row {idx+1}.')
        if r.phenotype not in direct_checked and r.Ti_s == 1.5 and r.RR_min == 25 and r.waveform == 'square':
            direct = np.zeros(len(C))
            for _ in range(10000):
                nxt = decay * inspiration(flow, C, R, direct).end_volume
                if np.max(np.abs(nxt-direct)) < PERIODIC_ATOL_L:
                    direct = nxt
                    break
                direct = nxt
            direct_error = float(np.max(np.abs(direct - Vee)))
            if direct_error > 1e-10:
                raise ValueError('Cycle-map and direct multi-breath solutions disagree.')
            max_direct_error = max(max_direct_error, direct_error)
            direct_checked.add(r.phenotype)
        Dexp = CONVERSION * np.sum((per.end_volume**2 - next_Vee**2) / (2.0 * C))
        cycle_elastic = CONVERSION * np.sum((next_Vee**2 - Vee**2) / (2.0 * C))
        # With forward Euler, W = D_insp + D_exp + delta_U_cycle - Euler_increment.
        adjusted_residual = (per.global_energy - per.resistive_energy - Dexp
                             - cycle_elastic + per.euler_energy_increment)
        if abs(adjusted_residual) > 1e-10:
            raise ValueError('Discrete energy-balance check failed.')
        if not np.isclose(np.sum(per.end_volume - next_Vee), r.VT_ml / 1000., atol=1e-10, rtol=0):
            raise ValueError('Periodic expired volume does not equal delivered tidal volume.')
        row = r.to_dict()
        # Legacy finite-end-inspiration values are kept out of the new export.
        row.pop('Pplat_cmH2O', None)
        row.pop('driving_pressure_cmH2O', None)
        row.update(scenario_id=idx+1, Te_s=te,
                   periodic_MP_J_min=per.global_energy * r.RR_min,
                   periodic_EII=per.eii, periodic_DCES=per.dces,
                   periodic_Ppeak_cmH2O=per.peak_above_peep + r.PEEP_cmH2O,
                   excess_end_expiratory_volume_mL=float(np.sum(Vee) * 1000.),
                   PEEPi_eq_cmH2O=float(np.sum(Vee) / np.sum(C)),
                   branch_PEEPi_min_cmH2O=float(np.min(Vee/C)),
                   branch_PEEPi_max_cmH2O=float(np.max(Vee/C)),
                   expiratory_resistive_energy_J=float(Dexp),
                   inspiratory_resistive_energy_J=per.resistive_energy,
                   inspiratory_work_J=per.global_energy,
                   euler_energy_increment_J=per.euler_energy_increment,
                   cycle_elastic_change_J=float(cycle_elastic),
                   cycle_map_iterations=iterations, cycle_closure_max_L=closure,
                   discrete_energy_balance_residual_J=float(adjusted_residual),
                   MP_increase_percent=100.0*(per.global_energy*r.RR_min/r.MP_J_min-1.0))
        for i in range(len(C)):
            row[f'periodic_REF_{i+1}'] = float(per.ref[i])
            row[f'periodic_positive_energy_J_{i+1}'] = float(per.positive_energy[i])
            row[f'periodic_signed_work_J_{i+1}'] = float(per.signed_energy[i])
            row[f'end_expiratory_volume_L_{i+1}'] = float(Vee[i])
            row[f'branch_PEEPi_cmH2O_{i+1}'] = float(Vee[i]/C[i])
        rows.append(row)
    output = add_periodic_scores(pd.DataFrame(rows))
    checks = self_tests()
    checks.update(baseline_max_absolute_difference=max_baseline_error,
                  direct_multibreath_max_volume_difference_L=max_direct_error,
                  direct_multibreath_configurations_checked=sorted(direct_checked),
                  max_cycle_closure_L=float(output.cycle_closure_max_L.max()),
                  max_discrete_energy_balance_residual_J=float(output.discrete_energy_balance_residual_J.abs().max()),
                  max_cycle_map_iterations=int(output.cycle_map_iterations.max()),
                  original_primary_count=len(output), periodic_feasible_count=int(output.periodic_feasible.sum()),
                  python_version=platform.python_version(), numpy_version=np.__version__, pandas_version=pd.__version__,
                  dt_s=DT, periodic_tolerance_L=PERIODIC_ATOL_L,
                  selection_rtol=SELECTION_RTOL, selection_atol=SELECTION_ATOL,
                  note='No time-step or parameter-grid convergence claim; checks validate this implementation only.')
    return output, checks


def export_table(df: pd.DataFrame, selected: pd.DataFrame, summary: pd.DataFrame, out: Path) -> Path:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.top_margin, sec.bottom_margin = Cm(1.35), Cm(1.35)
    sec.left_margin = sec.right_margin = Cm(1.4)
    norm = doc.styles['Normal']
    norm.font.name, norm.font.size = 'Arial', Pt(9)
    norm.font.color.rgb = RGBColor(0,0,0)
    norm.paragraph_format.space_after = Pt(5)
    norm.paragraph_format.line_spacing = 1.06
    def p(text: str, bold=False, size=9):
        para = doc.add_paragraph()
        para.paragraph_format.keep_with_next = bold
        run = para.add_run(text); run.bold=bold; run.font.size=Pt(size)
        return para
    def table(heads, rows, widths, size=8.5):
        t=doc.add_table(rows=1,cols=len(heads));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.style='Table Grid'
        for col,w in zip(t.columns,widths):col.width=Cm(w)
        for c,h,w in zip(t.rows[0].cells,heads,widths):c.text=h;c.width=Cm(w)
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells=t.add_row().cells
            for c,v,w in zip(cells,row,widths):c.text=str(v);c.width=Cm(w)
        for j,row in enumerate(t.rows):
            no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
            for i,c in enumerate(row.cells):
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for pa in c.paragraphs:
                    pa.paragraph_format.space_after=Pt(5);pa.paragraph_format.space_before=Pt(5)
                    pa.alignment=WD_ALIGN_PARAGRAPH.LEFT if i<2 else WD_ALIGN_PARAGRAPH.CENTER
                    for r in pa.runs:r.font.size=Pt(size);r.bold=(j==0)
        return t
    def rng(row,field,d=3):return f'{row[field+"_min"]:.{d}f}–{row[field+"_max"]:.{d}f}'
    p('Supplementary Table S3. Sensitivity to passive expiration and incomplete expiratory emptying.',True,12)
    p('A. Matched comparison across the original primary scenarios.',True,10)
    p('The additional periodic analysis retained the original 324 scenarios (81 per configuration). All remained below the 35 cmH₂O peak-pressure limit. Differences below are evaluated at identical ventilator settings; original results were not replaced.')
    rows=[]
    for _,s in summary.iterrows():
        rows.append([LABEL[s.phenotype],s.n_periodic_feasible,rng(s,'Te_s',1),rng(s,'MP_increase_percent',3),
                     rng(s,'excess_end_expiratory_volume_mL',3),rng(s,'PEEPi_eq_cmH2O',3)])
    table(['Configuration','Feasible n','Tₑ range (s)','MP increase (%)','Excess end-expiratory\nvolume (mL)','PEEPi,eq range\n(cmH₂O)'],rows,[4.3,2.0,3.0,4.0,6.7,6.9],9)
    p('B. Effect of periodic volume history on the original minimum-MP settings.',True,10)
    p('All four rows use Vₜ 4 mL/kg, RR 25/min, Tᵢ 1.5 s, square flow and no pause. These settings also remained the minimum-MP selections after re-evaluation at periodic steady state.')
    rows=[]
    for phe in ORDER:
        r=selected[(selected.phenotype==phe)&(selected.selection_basis=='original')&(selected.strategy=='minimum_MP')].iloc[0]
        rows.append([LABEL[phe],f'{r.MP_J_min:.3f}',f'{r.periodic_MP_J_min:.3f}',f'{r.MP_increase_percent:.2f}',
                     f'{r.EII:.3f}',f'{r.periodic_EII:.3f}',f'{r.DCES:.3f}',f'{r.periodic_DCES:.3f}',
                     f'{r.excess_end_expiratory_volume_mL:.2f}',f'{r.PEEPi_eq_cmH2O:.3f}'])
    table(['Configuration','Original\nMP','Periodic\nMP','MP increase\n(%)','Original\nEII','Periodic\nEII','Original\nDCES','Periodic\nDCES','Excess end-exp.\nvolume (mL)','PEEPi,eq\n(cmH₂O)'],rows,[3.9,2.0,2.0,2.6,2.0,2.0,2.0,2.0,4.5,3.9],8.5)
    p('MP values are in J/min. MP increase = 100 × (periodic MP / original MP − 1). The original model resets compartmental volumes to zero before inspiration; the additional model retains the volume remaining after finite passive expiration.')
    p('Vee,i denotes excess end-expiratory compartment volume above its static equilibrium volume at the prescribed PEEP. PEEPi,eq = ΣVee,i / ΣCᵢ is the analytical pressure increment above prescribed PEEP after hypothetical complete zero-flow equilibration. It is not a patient measurement and is distinct from the maintained airway-outlet pressure during expiration.')
    doc.add_page_break()
    p('Supplementary Table S3 (continued).',True,12)
    p('C. Strategies reselected using periodic-state outputs.',True,10)
    p('Minimum-MP and balanced selection were repeated separately within each configuration. MP, EII and DCES were min–max normalized within the periodically feasible set before calculating the balanced score. All selections were unique within the stated numerical tolerance.')
    rows=[]
    for phe in ORDER:
        for st in ('minimum_MP','balanced'):
            r=selected[(selected.phenotype==phe)&(selected.selection_basis=='periodic')&(selected.strategy==st)].iloc[0]
            rows.append([LABEL[phe],'Minimum MP' if st=='minimum_MP' else 'Balanced',int(r.VT_ml_per_kg),int(r.RR_min),f'{r.Ti_s:.1f}',r.waveform.capitalize(),f'{r.pause_fraction:.1f}',f'{r.periodic_MP_J_min:.3f}',f'{r.periodic_EII:.3f}',f'{r.periodic_DCES:.3f}',f'{r.excess_end_expiratory_volume_mL:.2f}',f'{r.PEEPi_eq_cmH2O:.3f}'])
    table(['Configuration','Selection','Vₜ\n(mL/kg)','RR\n(/min)','Tᵢ\n(s)','Waveform','Pause\nfraction','MP\n(J/min)','EII','DCES','Excess end-exp.\nvolume (mL)','PEEPi,eq\n(cmH₂O)'],rows,[3.5,2.7,1.65,1.25,1.2,2.8,1.7,1.8,1.65,1.65,2.9,2.1],8.4)
    p('All selected settings provide minute ventilation of 7.0 L/min at predicted body weight 70 kg and prescribed PEEP 5 cmH₂O. Balanced selection remained unchanged in the compliance-dominant configuration. In the other three configurations it changed from Vₜ 4 mL/kg and RR 25/min to Vₜ 5 mL/kg and RR 20/min, while retaining the originally selected Tᵢ, waveform and pause fraction.')
    p('Additional-model assumptions and numerical implementation',True,10)
    p('Expiration was passive into fixed PEEP, with expiratory resistance equal to inspiratory resistance in every branch. Tₑ = 60/RR − Tᵢ; Tᵢ includes the pause. Each branch emptied as Vᵢ(t) = Vᵢ(Tᵢ) exp[−(t−Tᵢ)/(RᵢCᵢ)]. The affine full-cycle map was iterated until the maximum between-cycle change in end-expiratory volume was below 10⁻¹² L per compartment; the final inspiration was then integrated explicitly. Original inspiratory waveforms, the 0.002-s Euler step, and energy definitions were retained. Selection ties used relative tolerance 10⁻¹⁰ and absolute tolerance 10⁻¹².')
    p('Global inspiratory MP and regional positive-inflow energies remained referenced to prescribed external PEEP, not total PEEP. Expiratory resistive dissipation, calculated from the decrease in stored elastic energy, is reported separately in the numerical data and was not added again to inspiratory work. Cycle closure, expired tidal volume, and the discrete energy balance including the Euler-method energy increment were checked; these checks are not a time-step convergence study.')
    p('Abbreviations: MP, global inspiratory mechanical power above prescribed PEEP; EII, Energy Inequality Index; DCES, Dominant Compartment Energy Share; RR, respiratory rate; Vₜ, tidal volume; Tᵢ, total inspiratory time; Tₑ, expiratory time. EII and DCES are dimensionless.')
    p('This sensitivity analysis does not represent expiratory flow limitation, airway closure, nonlinear or direction-dependent resistance, shared expiratory circuit resistance, patient effort, or recruitment–derecruitment. It demonstrates model-specific sensitivity to finite emptying and does not quantify auto-PEEP, injury or optimal settings in patients.')
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Supplementary Table S3 | ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    # Use real editable Word subscripts rather than underscore-style notation.
    import re
    from copy import deepcopy
    from docx.text.run import Run
    notation={'PEEPi,eq':('PEEP','i,eq'),'Vee,i':('V','ee,i'),'Vₜ':('V','T'),
              'Tᵢ':('T','i'),'Tₑ':('T','e'),'Vᵢ':('V','i'),'Rᵢ':('R','i'),'Cᵢ':('C','i')}
    pattern=re.compile('|'.join(re.escape(k) for k in sorted(notation,key=len,reverse=True)))
    paragraphs=list(doc.paragraphs)
    for tb in doc.tables:
        for row in tb.rows:
            for cell in row.cells:paragraphs.extend(cell.paragraphs)
    for para in paragraphs:
        for run in list(para.runs):
            text=run.text
            if not pattern.search(text):continue
            pieces=[];cursor=0
            for hit in pattern.finditer(text):
                pieces.append((text[cursor:hit.start()],False))
                base,sub=notation[hit.group()];pieces.extend([(base,False),(sub,True)])
                cursor=hit.end()
            pieces.append((text[cursor:],False))
            for content,issub in pieces:
                if not content:continue
                element=OxmlElement('w:r')
                if run._r.rPr is not None:element.append(deepcopy(run._r.rPr))
                run._r.addprevious(element);nr=Run(element,para);nr.text=content;nr.font.subscript=issub
            run._r.getparent().remove(run._r)
    doc.core_properties.title='Supplementary Table S3: Passive expiration sensitivity'
    doc.core_properties.author='Tomasz Urbankowski'
    path=out/'Supplementary_Table_S3_expiration.docx';doc.save(path);return path


def main() -> None:
    parser=argparse.ArgumentParser(description=__doc__)
    here=Path(__file__).resolve().parent
    parser.add_argument('--model-script',type=Path,default=here/MODEL_NAME)
    parser.add_argument('--input-csv',type=Path,default=here/INPUT_NAME)
    parser.add_argument('--output-dir',type=Path,default=here/'expiration_sensitivity')
    parser.add_argument('--no-docx',action='store_true')
    args=parser.parse_args()
    for path in (args.model_script,args.input_csv):
        if not path.is_file():parser.error(f'Input file does not exist: {path}')
    out=args.output_dir.resolve();out.mkdir(parents=True,exist_ok=True)
    df=pd.read_csv(args.input_csv);validate_input(df)
    result,checks=analyse(df,load_model(args.model_script.resolve()))
    selected=select_strategies(result);summary=configuration_summary(result,selected)
    result.to_csv(out/'expiration_primary_scenarios.csv',index=False)
    selected.to_csv(out/'expiration_selected_strategies.csv',index=False)
    summary.to_csv(out/'expiration_summary.csv',index=False)
    for name,path in [('input',args.input_csv),('model',args.model_script)]:
        checks[name+'_filename']=path.name
        checks[name+'_sha256']=hashlib.sha256(path.read_bytes()).hexdigest()
    (out/'expiration_verification.json').write_text(json.dumps(checks,indent=2),encoding='utf-8')
    if not args.no_docx:
        try:print(export_table(result,selected,summary,out))
        except ImportError as exc:raise SystemExit('Install python-docx or use --no-docx; CSV outputs were saved.') from exc
    print(summary.to_string(index=False))
    print('Verification:',json.dumps(checks,indent=2))
    print('Output directory:',out)


if __name__=='__main__':main()
