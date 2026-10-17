#!/usr/bin/env python3
"""Reproduce the R1.5 within-model comparison with simpler global indices.

No change is made to the uploaded model's mechanics, integration, flow waveforms,
regional-energy definition, feasibility rules, or balanced score.

Run from the directory containing the original simulation script:
    python compare_simple_indices_R1_5.py
Or reuse an already generated primary feasible CSV:
    python compare_simple_indices_R1_5.py --input-csv path/to/primary.csv

Dependencies: numpy, pandas, python-docx (unless --no-docx is used).
DP_eq = VT / sum(C_i) is an analytical, fully equilibrated, zero-flow reference.
It is NOT the pressure at the end of a finite pause and NOT transpulmonary pressure.
The original columns labelled Pplat and driving_pressure are NOT used here.
4DPRR = 4 * DP_eq + RR uses numerical inputs in cmH2O and breaths/min; this is a
unit-convention-based score, not mechanical power or a physical energy quantity.
All comparisons are descriptive. No inference about injury or mortality is made.
"""
from __future__ import annotations

import argparse
from contextlib import contextmanager
import hashlib
import importlib.util
import json
import os
from pathlib import Path
from typing import Iterator

import numpy as np
import pandas as pd

ORDER = ('compliance_dominant', 'resistance_dominant', 'mixed', 'severe_mixed')
LABELS = {'compliance_dominant': 'Compliance-dominant',
          'resistance_dominant': 'Resistance-dominant',
          'mixed': 'Mixed', 'severe_mixed': 'Severe mixed'}
OBJECTIVES = {
    'minimum_MP': 'MP_J_min',
    'minimum_Ppeak': 'Ppeak_cmH2O',
    'balanced': 'balanced_score',
    'minimum_DP_eq': 'DP_eq_cmH2O',
    'minimum_4DPRR': 'index_4DPRR',
}
STRATEGY_LABELS = {'minimum_MP': 'Minimum MP', 'minimum_Ppeak': 'Minimum Ppeak',
                   'balanced': 'Balanced', 'minimum_DP_eq': 'Minimum DP_eq',
                   'minimum_4DPRR': 'Minimum 4DPRR'}
RTOL, ATOL = 1e-10, 1e-12


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


@contextmanager
def working_directory(path: Path) -> Iterator[None]:
    previous = Path.cwd()
    path.mkdir(parents=True, exist_ok=True)
    try:
        os.chdir(path)
        yield
    finally:
        os.chdir(previous)


def run_model(model_path: Path, output: Path) -> pd.DataFrame:
    spec = importlib.util.spec_from_file_location('uploaded_ventilation_model', model_path)
    if spec is None or spec.loader is None:
        raise ImportError(f'Cannot load model: {model_path}')
    model = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(model)
    if not hasattr(model, 'run_parameter_sweep'):
        raise ValueError('The model must expose run_parameter_sweep().')
    with working_directory(output / 'original_results'):
        results = model.run_parameter_sweep()
    return results['primary_subset'].copy()


def prepare(df: pd.DataFrame) -> pd.DataFrame:
    required = ['phenotype', 'VT_ml', 'VT_ml_per_kg', 'RR_min', 'Ti_s', 'waveform',
                'pause_fraction', 'PEEP_cmH2O', 'C_total_L_per_cmH2O',
                'minute_ventilation_L_min', 'Ppeak_cmH2O', 'MP_J_min',
                'EII', 'DCES', 'balanced_score']
    missing = sorted(set(required) - set(df.columns))
    if missing:
        raise ValueError(f'Missing columns: {missing}')
    x = df.copy().reset_index(drop=True)
    x.insert(0, 'scenario_id', np.arange(1, len(x) + 1))
    if not x[required].notna().all().all():
        raise ValueError('Required columns contain missing values.')
    if len(x) != 324 or x.groupby('phenotype').size().to_dict() != dict.fromkeys(ORDER, 81):
        raise ValueError('Expected the original primary analysis: 324 rows, 81/configuration.')
    if not ((x.Ppeak_cmH2O <= 35.0 + ATOL) &
            x.minute_ventilation_L_min.between(6.5 - ATOL, 7.5 + ATOL) &
            np.isclose(x.PEEP_cmH2O, 5.0)).all():
        raise ValueError('Input does not match the original primary feasibility conditions.')
    if not (x.C_total_L_per_cmH2O > 0).all():
        raise ValueError('Total compliance must be positive.')
    # These names make clear what the legacy columns actually represent. The
    # source CSV remains untouched, and neither column is used in comparisons.
    x = x.rename(columns={
        'Pplat_cmH2O': 'legacy_end_inspiratory_sample_cmH2O',
        'driving_pressure_cmH2O': 'legacy_end_inspiratory_sample_minus_PEEP_cmH2O'})
    x['DP_eq_cmH2O'] = (x.VT_ml / 1000.0) / x.C_total_L_per_cmH2O
    x['index_4DPRR'] = 4.0 * x.DP_eq_cmH2O + x.RR_min
    return x


def compare(x: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    summary, singleton, minima = [], [], []
    metrics = ['DP_eq_cmH2O', 'index_4DPRR', 'Ppeak_cmH2O', 'MP_J_min', 'EII', 'DCES']
    for phenotype in ORDER:
        sub = x[x.phenotype.eq(phenotype)]
        for strategy, objective in OBJECTIVES.items():
            minimum = float(sub[objective].min())
            tied = sub[np.isclose(sub[objective], minimum, rtol=RTOL, atol=ATOL)].copy()
            tied.insert(1, 'comparison_strategy', strategy)
            minima.append(tied)
            record = {'phenotype': phenotype, 'strategy': strategy,
                      'objective': objective, 'objective_minimum': minimum,
                      'n_minimizers': len(tied)}
            for col in ['VT_ml_per_kg', 'RR_min', 'Ti_s', 'waveform', 'pause_fraction']:
                record[col + '_values'] = ';'.join(str(v) for v in sorted(tied[col].unique()))
            for col in metrics:
                record[col + '_min'] = float(tied[col].min())
                record[col + '_max'] = float(tied[col].max())
            summary.append(record)
            if len(tied) == 1:
                singleton.append(tied)
    return (pd.DataFrame(summary), pd.concat(singleton, ignore_index=True),
            pd.concat(minima, ignore_index=True))


def validate_against_table1(single: pd.DataFrame, summary: pd.DataFrame) -> None:
    # Values below are the rounded published-input values in the attached Table 1.
    expected = {
        ('compliance_dominant', 'minimum_MP'): (3.52, .306, .237, 13.66),
        ('compliance_dominant', 'balanced'): (5.14, .229, .215, 16.09),
        ('resistance_dominant', 'minimum_MP'): (3.63, .067, .181, 13.85),
        ('resistance_dominant', 'balanced'): (3.86, .025, .171, 13.11),
        ('mixed', 'minimum_MP'): (3.32, .376, .264, 13.14),
        ('mixed', 'balanced'): (3.57, .364, .256, 12.41),
        ('severe_mixed', 'minimum_MP'): (2.82, .569, .324, 11.89),
        ('severe_mixed', 'balanced'): (3.04, .558, .316, 11.28),
    }
    for (phenotype, strategy), targets in expected.items():
        rows = single[single.phenotype.eq(phenotype) & single.comparison_strategy.eq(strategy)]
        if len(rows) != 1:
            raise ValueError(f'Expected a unique baseline selection for {phenotype}/{strategy}.')
        values = rows.iloc[0][['MP_J_min','EII','DCES','Ppeak_cmH2O']].to_numpy(float)
        if not np.all(np.abs(values - targets) <= np.array([.00501,.000501,.000501,.00501])):
            raise ValueError(f'Reproduced values do not match attached Table 1: {phenotype}/{strategy}.')
    # These assertions document observed results of this specific model/grid,
    # not assumptions about other possible models or later revisions.
    for strategy in ['minimum_DP_eq', 'minimum_4DPRR']:
        s = summary[summary.strategy.eq(strategy)]
        if not s.n_minimizers.eq(27).all():
            raise ValueError('Minimizer counts differ from the original R1.5 comparison.')


def export_docx(single: pd.DataFrame, summary: pd.DataFrame, output: Path) -> Path:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.top_margin = sec.bottom_margin = Cm(1.45)
    sec.left_margin = sec.right_margin = Cm(1.5)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.06

    def paragraph(text: str, bold: bool = False, size: float = 10) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = bold
        r = p.add_run(text)
        r.bold, r.font.size = bold, Pt(size)

    def table(headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        t.style = 'Table Grid'
        for c, width in zip(t.columns, widths):
            c.width = Cm(width)
        for cell, text, width in zip(t.rows[0].cells, headers, widths):
            cell.width = Cm(width)
            cell.text = text
        hdr = OxmlElement('w:tblHeader')
        t.rows[0]._tr.get_or_add_trPr().append(hdr)
        for values in rows:
            row = t.add_row()
            for c, text, width in zip(row.cells, values, widths):
                c.width = Cm(width)
                c.text = text
        for ri, row in enumerate(t.rows):
            trpr = row._tr.get_or_add_trPr()
            trpr.append(OxmlElement('w:cantSplit'))
            for ci, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci < 2 else WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.line_spacing = 1.0
                    for r in p.runs:
                        r.font.name = 'Arial'
                        r.font.size = Pt(9.5)
                        r.bold = ri == 0
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    paragraph('Supplementary Table S2. Comparison with simpler global indices in the primary analysis.', True, 12)
    paragraph('A. Settings and model outputs for minimum-MP, minimum-peak-pressure, and balanced selections.', True)
    paragraph('All selections below used tidal volume 4 mL/kg predicted body weight (280 mL), respiratory rate 25 breaths/min, minute ventilation 7.0 L/min, and PEEP 5 cmH₂O.')
    rows = []
    for phen in ORDER:
        for strat in ['minimum_MP', 'minimum_Ppeak', 'balanced']:
            r = single[single.phenotype.eq(phen) & single.comparison_strategy.eq(strat)].iloc[0]
            rows.append([LABELS[phen], STRATEGY_LABELS[strat], f'{r.Ti_s:.1f}',
                         str(r.waveform).capitalize(), f'{r.pause_fraction:.1f}',
                         f'{r.DP_eq_cmH2O:.3f}', f'{r.index_4DPRR:.3f}',
                         f'{r.Ppeak_cmH2O:.2f}', f'{r.MP_J_min:.2f}',
                         f'{r.EII:.3f}', f'{r.DCES:.3f}'])
    table(['Configuration', 'Selection', 'Ti\n(s)', 'Flow waveform', 'Pause\nfraction',
           'ΔP_eq\n(cmH₂O)', '4DPRR\n(score)', 'Ppeak\n(cmH₂O)',
           'MP\n(J/min)', 'EII', 'DCES'], rows,
          [3.9, 3.2, 1.3, 3.0, 2.0, 2.2, 2.1, 2.1, 1.7, 1.6, 1.6])
    paragraph('ΔP_eq = VT/Ctotal, with VT in L and Ctotal = ΣCi in L/cmH₂O, is the analytical airway-pressure increment above PEEP after complete zero-flow equilibration of the model. It is not an end-of-inspiration sample, a finite-pause pressure, or a patient-specific transpulmonary pressure.')
    paragraph('4DPRR = 4 × ΔP_eq + RR is evaluated using numerical inputs in cmH₂O and breaths/min. This composite score follows the form described by Costa et al. (2021); it is not mechanical power. The static driving-pressure formulation follows Amato et al. (2015).')
    paragraph('Ti, total inspiratory time; pause fraction, fraction of Ti; Ppeak, peak airway-opening pressure; MP, global inspiratory mechanical power above PEEP; EII, Energy Inequality Index; DCES, Dominant Compartment Energy Share. EII and DCES are dimensionless. The regional-energy definition and all original simulation settings were retained.')

    doc.add_page_break()
    paragraph('Supplementary Table S2 (continued).', True, 12)
    paragraph('B. Tied minima of equilibrium driving pressure and the 4DPRR index.', True)
    paragraph('In each configuration, the two criteria had exactly the same set of 27 minimizers: VT 4 mL/kg, RR 25 breaths/min, all three inspiratory times, all three flow waveforms, and all three pause fractions. No arbitrary single setting was selected from a tied set. The ranges below describe all minimizers, including both the minimum-MP and balanced selections.')
    rows = []
    for phen in ORDER:
        r = summary[summary.phenotype.eq(phen) & summary.strategy.eq('minimum_DP_eq')].iloc[0]
        def rng(col: str, precision: int) -> str:
            return f'{r[col+"_min"]:.{precision}f}–{r[col+"_max"]:.{precision}f}'
        rows.append([LABELS[phen], str(int(r.n_minimizers)), f'{r.DP_eq_cmH2O_min:.3f}',
                     f'{r.index_4DPRR_min:.3f}', rng('MP_J_min',2),
                     rng('Ppeak_cmH2O',2), rng('EII',3), rng('DCES',3)])
    table(['Configuration', 'Minimizers\n(n)', 'Minimum ΔP_eq\n(cmH₂O)',
           'Minimum 4DPRR\n(score)', 'MP range\n(J/min)', 'Ppeak range\n(cmH₂O)',
           'EII range', 'DCES range'], rows,
          [4.0, 2.4, 3.2, 3.2, 3.4, 3.4, 3.0, 3.0])
    paragraph('The comparisons were performed separately within each configuration among the same 81 primary feasible scenarios (324 in total), using minute ventilation 6.5–7.5 L/min and Ppeak ≤35 cmH₂O. Ties were identified with relative tolerance 10⁻¹⁰ and absolute tolerance 10⁻¹². Minimum-peak-pressure selections were unique in this dataset; all used a linearly decelerating waveform, Ti 1.5 s, and no pause.')
    paragraph('These descriptive comparisons assess whether scalar global criteria distinguish model-derived energy distributions. They do not assess prediction of VILI or mortality and do not establish the clinical superiority of any selection strategy.')
    paragraph('References', True)
    paragraph('Amato MBP et al. Driving pressure and survival in the acute respiratory distress syndrome. N Engl J Med. 2015;372:747–755. doi:10.1056/NEJMsa1410639.', size=9)
    paragraph('Costa ELV et al. Ventilatory variables and mechanical power in patients with acute respiratory distress syndrome. Am J Respir Crit Care Med. 2021;204:303–311. doi:10.1164/rccm.202009-3467OC.', size=9)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run('Supplementary Table S2 | ')
    run.font.size = Pt(8)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    footer._p.append(fld)
    # Use real Word subscripts rather than underscore notation in the table.
    from copy import deepcopy
    import re
    from docx.text.run import Run
    notation = {'ΔP_eq': ('ΔP', 'eq'), 'Ppeak': ('P', 'peak'),
                'Ctotal': ('C', 'total'), 'ΣCi': ('ΣC', 'i'),
                'VT': ('V', 'T'), 'Ti': ('T', 'i')}
    pattern = re.compile(r'(?<![A-Za-z])(' + '|'.join(map(re.escape, notation)) + r')(?![A-Za-z])')
    all_paragraphs = list(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                all_paragraphs.extend(cell.paragraphs)
    for p in all_paragraphs:
        for run in list(p.runs):
            text = run.text
            if not pattern.search(text):
                continue
            pieces, cursor = [], 0
            for match in pattern.finditer(text):
                pieces.append((text[cursor:match.start()], False))
                base, subscript = notation[match.group()]
                pieces.extend([(base, False), (subscript, True)])
                cursor = match.end()
            pieces.append((text[cursor:], False))
            for content, is_subscript in pieces:
                if not content:
                    continue
                element = OxmlElement('w:r')
                if run._r.rPr is not None:
                    element.append(deepcopy(run._r.rPr))
                run._r.addprevious(element)
                new_run = Run(element, p)
                new_run.text = content
                new_run.font.subscript = is_subscript
            run._r.getparent().remove(run._r)
    doc.core_properties.title = 'Supplementary Table S2: Comparison with simpler global indices'
    doc.core_properties.author = 'Tomasz Urbankowski'
    result = output / 'Supplementary_Table_S2_simple_indices.docx'
    doc.save(result)
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input-csv', type=Path, help='Original primary feasible results CSV.')
    parser.add_argument('--model-script', type=Path, help='Original uploaded simulation script.')
    parser.add_argument('--output-dir', type=Path, default=Path(__file__).resolve().parent / 'R1_5_comparison')
    parser.add_argument('--no-docx', action='store_true')
    args = parser.parse_args()
    out = args.output_dir.resolve()
    out.mkdir(parents=True, exist_ok=True)
    model_path = None
    if args.input_csv:
        input_path = args.input_csv.resolve()
        if not input_path.is_file():
            parser.error(f'Input file does not exist: {input_path}')
        raw = pd.read_csv(input_path)
    else:
        if args.model_script:
            model_path = args.model_script.resolve()
        else:
            candidates = sorted(Path(__file__).resolve().parent.glob('multiobjective_ventilation_optimization_scaled_CR_MVtarget*.py'))
            if len(candidates) != 1:
                parser.error('Supply --model-script; exactly one original model script could not be identified.')
            model_path = candidates[0]
        if not model_path.is_file():
            parser.error(f'Model file does not exist: {model_path}')
        raw = run_model(model_path, out)
        input_path = out / 'original_results' / 'MVtarget_primary_feasible_simulation_results.csv'
    x = prepare(raw)
    summary, single, minima = compare(x)
    validate_against_table1(single, summary)
    x.to_csv(out / 'R1_5_primary_with_simple_indices.csv', index=False)
    summary.to_csv(out / 'R1_5_minimizer_summary.csv', index=False)
    single.to_csv(out / 'R1_5_unique_selected_strategies.csv', index=False)
    minima.to_csv(out / 'R1_5_all_minimizers.csv', index=False)
    verification = {
        'input_csv': str(input_path), 'input_sha256': sha256(input_path),
        'model_path': str(model_path) if model_path else None,
        'model_sha256': sha256(model_path) if model_path else None,
        'primary_rows': len(x), 'rows_per_configuration': x.groupby('phenotype').size().to_dict(),
        'original_Table1_reproduced_at_reported_precision': True,
        'rtol': RTOL, 'atol': ATOL,
        'legacy_Pplat_and_driving_pressure_columns_used': False,
        'clinical_validation_performed': False,
        'numpy_version': np.__version__, 'pandas_version': pd.__version__,
    }
    (out / 'verification.json').write_text(json.dumps(verification, indent=2), encoding='utf-8')
    if not args.no_docx:
        try:
            print(export_docx(single, summary, out))
        except ImportError as exc:
            raise SystemExit('Install python-docx or use --no-docx. CSV outputs have been saved.') from exc
    print(f'Primary scenarios: {len(x)}. Original Table 1 reproduced at reported precision.')
    print(summary[['phenotype','strategy','n_minimizers','objective_minimum']].to_string(index=False))
    print(f'Outputs: {out}')


if __name__ == '__main__':
    main()
